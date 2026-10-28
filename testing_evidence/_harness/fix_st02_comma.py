from docx import Document

path = r"C:\Users\HomePC\Downloads\FRAMS_REVISED_DRAFT_UNIT_INTEGRATION_SYSTEM_PASSED.docx"
old = "\u201cBorrower account approved successfully,\u201d"
new = "\u201cBorrower account approved successfully.\u201d"
doc = Document(path)
hits = 0
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            if old in cell.text:
                for paragraph in cell.paragraphs:
                    if old in paragraph.text:
                        if paragraph.runs:
                            paragraph.runs[0].text = paragraph.text.replace(old, new)
                            for run in paragraph.runs[1:]:
                                run.text = ""
                            hits += 1
doc.save(path)
print("hits", hits)
