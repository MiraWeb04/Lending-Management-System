from docx import Document

path = r"C:\Users\HomePC\Downloads\FRAMS_REVISED_DRAFT_UNIT_INTEGRATION_SYSTEM_PASSED.docx"
doc = Document(path)

print("=== BODY MATCHES ===")
for i, paragraph in enumerate(doc.paragraphs):
    text = paragraph.text
    if (
        "Figure B10" in text
        or "Figure B11" in text
        or "Figure B12Evidence" in text
        or "Figure B13Evidence" in text
        or "Not Yet" in text
        or "not yet tested" in text.lower()
        or "12 system" in text
        or text.startswith("The results of unit")
        or "ST-01 through ST-12" in text
        or text.startswith("*Verified")
        or text.startswith("All 49 functional cases")
    ):
        print(i, repr(text[:500]))
        print("---")

print("=== SUMMARY TABLES ===")
for table in doc.tables:
    header = table.rows[0].cells[0].text.strip().lower()
    if header == "testing level":
        for row in table.rows:
            print(" | ".join(cell.text.replace("\n", " ")[:80] for cell in row.cells))
        print("====")

print("=== SAMPLE CASES ===")
for table in doc.tables:
    for row in table.rows:
        key = row.cells[0].text.strip()
        if key in ("UT-05", "IT-02", "IT-05", "ST-04", "ST-09", "PM-01", "PT-01"):
            print(key, "||", " || ".join(cell.text.replace("\n", " ")[:110] for cell in row.cells))
            print()
