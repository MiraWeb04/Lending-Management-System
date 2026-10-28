#!/usr/bin/env python3
"""Word document for penetration-test figures B38-B49."""
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

SHOTS = Path(r"C:\xampp\htdocs\lending_system\testing_evidence\05_Penetration_Testing\screenshots")
OUT = Path(r"C:\xampp\htdocs\lending_system\testing_evidence\05_Penetration_Testing\FRAMS_Penetration_Figures_B38_B49.docx")

CASES = [
    {
        "id": "PT-01",
        "figure": "Figure B38",
        "scenario": "Open the administrator dashboard, collector dashboard, payments, reports, and borrower pages while signed out.",
        "expected": "Staff pages redirect to the staff sign-in page. Borrower pages redirect to the borrower sign-in page. No records are shown.",
        "actual": "Signed-out requests for dashboard_lending.php, collector_dashboard_lending.php, payments_lending.php, and reports_lending.php each returned HTTP 302 to login_lending.php. The browser showed the staff sign-in page and no loan, client, or user records. Signed-out requests for borrower_dashboard_lending.php, borrower_my_loan_lending.php, and borrower_payment_history_lending.php each returned HTTP 302 to borrower_login_lending.php. The browser showed the borrower sign-in page.",
        "status": "PASS",
        "images": [
            ("B38_pt01_admin_dashboard.png", "Staff sign-in after opening the administrator dashboard while signed out."),
            ("B38_pt01_borrower_dashboard.png", "Borrower sign-in after opening the borrower dashboard while signed out."),
        ],
    },
    {
        "id": "PT-02",
        "figure": "Figure B39",
        "scenario": "Sign in with a wrong password and with an inactive account.",
        "expected": "No protected page opens. The inactive trial uses the same staff message as a wrong password, and the account status supports that trial.",
        "actual": "frams_test_admin with a wrong password stayed on the staff sign-in page and showed \"Invalid username or password.\" frams_test_inactive, whose users.status is Inactive, with the ordinary test password stayed on the same page and showed the same message. Neither attempt opened a dashboard.",
        "status": "PASS",
        "images": [
            ("B39_pt02_wrong_password.png", "Wrong password for the active administrator."),
            ("B39_pt02_inactive_account.png", "Inactive administrator. The message matches the wrong-password trial."),
        ],
    },
    {
        "id": "PT-03",
        "figure": "Figure B40",
        "scenario": "A borrower opens the users, reports, loan-release, expenses, and payments pages.",
        "expected": "Users, reports, loan release, and expenses return the borrower to the borrower dashboard. The payments page returns the borrower to that borrower payment history.",
        "actual": "Signed in as frams_test_borrower. users_lending.php, reports_lending.php, loan_release_lending.php, and expenses_lending.php each returned HTTP 302 to borrower_dashboard_lending.php. The browser showed the borrower dashboard, not user administration, reports, loan release, or expenses. payments_lending.php returned HTTP 302 to borrower_payment_history_lending.php. That history listed only this borrower's receipt RJRR-20261018-000004.",
        "status": "PASS",
        "images": [
            ("B40_pt03_users_to_borrower_dashboard.png", "Borrower dashboard after requesting user management."),
            ("B40_pt03_payments_to_history.png", "Borrower payment history after requesting the staff payments page."),
        ],
    },
    {
        "id": "PT-04",
        "figure": "Figure B41",
        "scenario": "A collector opens users, reports, borrower approval, and loan release.",
        "expected": "Each request returns to the collector dashboard, and no administrator record is changed.",
        "actual": "Signed in as frams_test_collector. users_lending.php, reports_lending.php, borrower_approval_lending.php, and loan_release_lending.php each returned HTTP 302 to collector_dashboard_lending.php. The browser showed the collector dashboard. No user, approval, or release form was shown, and the user count stayed at 6.",
        "status": "PASS",
        "images": [
            ("B41_pt04_users_to_collector_dashboard.png", "Collector dashboard after requesting user management. Reports, borrower approval, and loan release returned to this same dashboard."),
        ],
    },
    {
        "id": "PT-05",
        "figure": "Figure B42",
        "scenario": "Submit registration, payment recording, expense recording, user management, or loan release without the security token.",
        "expected": "The request ends with HTTP 419 and the message that the session is invalid or expired. No record is saved.",
        "actual": "Registration, payment recording, expense recording, user management, and loan release were each posted without the security token. Borrower approval was not used. All five responses were HTTP 419 with the heading \"Invalid or expired session\" and the text \"Please refresh the page and try again.\" Users stayed at 6, payments at 4, expenses at 0, and loan releases at 3. The probe username was not created.",
        "status": "PASS",
        "images": [
            ("B42_pt05_missing_token.png", "Registration submitted without the security token. The other four posts returned this same page with HTTP 419."),
        ],
    },
    {
        "id": "PT-06",
        "figure": "Figure B43",
        "scenario": "Enter a quote-style probe in sign-in and in a search box.",
        "expected": "Sign-in fails. Search does not reveal additional accounts, loans, or payments.",
        "actual": "The staff sign-in username frams_test_admin' OR '1'='1 was rejected with \"Invalid username or password\" and stayed on the sign-in page. The same quote in user search returned \"Showing 0 of 0 users\" and \"No users found.\" It did not list the administrator, collector, or borrower accounts.",
        "status": "PASS",
        "images": [
            ("B43_pt06_quote_login.png", "Quote-style sign-in probe rejected."),
            ("B43_pt06_quote_search.png", "Quote-style user search returned no accounts."),
        ],
    },
    {
        "id": "PT-07",
        "figure": "Figure B44",
        "scenario": "Enter a harmless markup marker in an encoded search field and display it.",
        "expected": "The marker is shown as ordinary text.",
        "actual": "The user-search value <b>pt07marker</b> was returned in the page as encoded text. The search box showed the marker as ordinary characters, including the angle brackets. It was not rendered as bold text, and the search still listed 0 users. The payments flash message was not used.",
        "status": "PASS",
        "images": [
            ("B44_pt07_marker_as_text.png", "The markup marker remains visible as text in the search box."),
        ],
    },
    {
        "id": "PT-08",
        "figure": "Figure B45",
        "scenario": "A borrower requests another borrower receipt address and document address.",
        "expected": "The receipt is refused or redirected, and the document request returns Access denied.",
        "actual": "Signed in as frams_test_borrower. Requesting payments_lending.php?view=44, the receipt for Borrower B, returned HTTP 302 to borrower_payment_history_lending.php. The history still showed only RJRR-20261018-000004 and did not show RJRR-20261018-000002. Requesting download_loan_document_lending.php for a document path that does not belong to this borrower returned HTTP 403 and the text \"Access denied.\"",
        "status": "PASS",
        "images": [
            ("B45_pt08_other_receipt_refused.png", "Other-borrower receipt request returned to this borrower's own payment history."),
            ("B45_pt08_document_access_denied.png", "Other-borrower document request. The response text is Access denied."),
        ],
    },
    {
        "id": "PT-09",
        "figure": "Figure B46",
        "scenario": "A collector opens a client, loan, or payment assigned to another collector.",
        "expected": "The unassigned record is not opened for collection.",
        "actual": "Signed in as frams_test_collector. Client 21, which is assigned to Collector B, returned HTTP 403: \"You do not have permission to view this client.\" Loan 30, FRAMS-TEST-B, returned HTTP 403: \"You do not have permission to view this loan.\" Payment 44 returned HTTP 403: \"You do not have access to this receipt.\" None of the three records opened.",
        "status": "PASS",
        "images": [
            ("B46_pt09_other_client.png", "Collector A denied client 21."),
            ("B46_pt09_other_loan.png", "Collector A denied loan 30."),
            ("B46_pt09_other_payment.png", "Collector A denied payment 44."),
        ],
    },
    {
        "id": "PT-10",
        "figure": "Figure B47",
        "scenario": "Upload a file that is not JPG, PNG, or PDF, and a file larger than 8 MB.",
        "expected": "The file is rejected and is not stored. A PDF is an allowed file and was not used as the invalid sample.",
        "actual": "On loan-application step 4, pt10-probe.txt was checked and rejected with \"Allowed file types: JPG, PNG, PDF.\" A file of 8,389,633 bytes was rejected with \"File must be 8 MB or smaller.\" Neither file was written to the document folder. The wizard does not keep that sentence on screen: it stays on step 4 and shows \"Required documents missing,\" with every document still marked \"Not uploaded yet.\"",
        "status": "PASS",
        "images": [
            ("B47_pt10_upload_rejected.png", "After the invalid uploads, step 4 still shows every document as not uploaded."),
        ],
    },
    {
        "id": "PT-11",
        "figure": "Figure B48",
        "scenario": "After signing out, reload a protected page.",
        "expected": "The reloaded page redirects to the sign-in page.",
        "actual": "The collector signed out, which opened the staff sign-in page. Opening collector_dashboard_lending.php again returned HTTP 302 to login_lending.php. The browser showed the staff sign-in page, so the protected dashboard required authentication again.",
        "status": "PASS",
        "images": [
            ("B48_pt11_after_logout.png", "Protected collector dashboard reloaded after sign-out."),
        ],
    },
    {
        "id": "PT-12",
        "figure": "Figure B49",
        "scenario": "Submit an incorrect, expired, and already used six-digit password-reset code.",
        "expected": "The password is not changed by an invalid code. The live reset flow uses a six-digit code that expires after ten minutes, allows five incorrect attempts, and is removed after a successful change. It does not issue a reset link.",
        "actual": "The forgot-password page asks for a username and says it will email a 6-digit code. It does not offer a reset link. The code form says the code expires in 10 minutes. A new code record was measured at 600 seconds. An incorrect code returned \"That code is incorrect or has expired.\" After five incorrect attempts, the next try returned \"Too many incorrect codes. Request a new code.\" and the code record was removed. A code whose expiry was already past returned the incorrect-or-expired message. A correct code allowed one password change, and that same code was then rejected. The borrower test password was restored afterward, and no reset row remained.",
        "status": "PASS",
        "images": [
            ("B49_pt12_reset_requests_code.png", "The reset page sends a 6-digit code and does not issue a link."),
            ("B49_pt12_code_form.png", "The code form states that the code expires in 10 minutes."),
            ("B49_pt12_incorrect_code.png", "An incorrect code does not change the password."),
            ("B49_pt12_too_many_codes.png", "The attempt after five incorrect codes is refused."),
            ("B49_pt12_expired_code.png", "An expired code is refused."),
            ("B49_pt12_password_changed.png", "A valid code can be used once."),
            ("B49_pt12_used_code.png", "The same code is refused after the password change."),
        ],
    },
]


def add_picture(document, filename):
    path = SHOTS / filename
    if not path.exists():
        document.add_paragraph("Missing screenshot: " + filename)
        return
    document.add_picture(str(path), width=Inches(6.4))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def main():
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    title = document.add_heading("FRAMS penetration test figures", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("Figures B38 to B49")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(
        "Execution date: 19 October 2026. Environment: isolated FRAMS test server at "
        "http://127.0.0.1:8099 and database frams_capstone_test. The existing "
        "lending_management database was not modified."
    )
    document.add_paragraph(
        "Results: 12 passed, 0 failed, 0 blocked, 0 not executed."
    )

    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for index, header in enumerate(["Test ID", "Figure", "Status", "Actual result"]):
        table.rows[0].cells[index].text = header
    for case in CASES:
        cells = table.add_row().cells
        cells[0].text = case["id"]
        cells[1].text = case["figure"]
        cells[2].text = case["status"]
        cells[3].text = case["actual"]
    document.add_page_break()

    for case in CASES:
        document.add_heading(case["figure"] + " — " + case["id"], level=1)
        document.add_paragraph("Scenario: " + case["scenario"])
        document.add_paragraph("Expected result: " + case["expected"])
        document.add_paragraph("Actual result: " + case["actual"])
        status = document.add_paragraph()
        run = status.add_run("Status: " + case["status"])
        run.bold = True
        run.font.color.rgb = RGBColor(0x16, 0x65, 0x34)
        for filename, caption_text in case["images"]:
            add_picture(document, filename)
            caption = document.add_paragraph(case["figure"] + ". " + caption_text)
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.runs[0].italic = True
        document.add_paragraph("")

    document.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
