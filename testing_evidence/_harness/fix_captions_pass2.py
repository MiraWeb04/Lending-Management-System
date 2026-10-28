from docx import Document

path = r"C:\Users\HomePC\Downloads\FRAMS_REVISED_DRAFT_UNIT_INTEGRATION_SYSTEM_PASSED.docx"
out_log = r"C:\xampp\htdocs\lending_system\testing_evidence\_harness\caption_pass2.txt"
doc = Document(path)

EN = "\u2013"
replacements = {
    "Figure B10": "Figure B10\nEvidence for IT-04: Payment Recording and Remaining Balance. The paid installment and remaining balance are shown in Figure B22.",
    "Figure B11": "Figure B11\nEvidence for IT-05: Loan Application and Eligibility Assessment",
    "Figure B12": "Figure B12\nEvidence for IT-06: Reports for the Payment Period",
    "Figure B13": "Figure B13\nEvidence for IT-07: Borrower Username on the Staff Sign-In",
}

lines = []
for paragraph in doc.paragraphs:
    text = paragraph.text
    key = text.split("\n", 1)[0].strip()
    if key in replacements and text.startswith("Figure B1"):
        # Only the evidence captions, which continue with "Evidence for".
        if "Evidence for" in text:
            if paragraph.runs:
                paragraph.runs[0].text = replacements[key]
                for run in paragraph.runs[1:]:
                    run.text = ""
            lines.append("UPDATED " + key)
    if text.startswith("The results of unit"):
        lines.append("PRESENTATION END: " + text[-700:])
        if "Not Yet Tested" in text or "retained without changes" in text:
            old = "The 12 system test cases remain Not Yet Tested in this report. Unit and integration testing entries are retained without changes."
            new = "Unit testing passed in 6 of 6 cases, integration testing passed in 7 of 7 cases, and system testing passed in 12 of 12 cases."
            if old in text:
                updated = text.replace(old, new)
            else:
                # Keep the PayMongo and penetration sentences and drop the not-yet closing.
                marker = "The 12 system test cases remain"
                updated = text[: text.find(marker)].rstrip() + " " + new
            if paragraph.runs:
                paragraph.runs[0].text = updated
                for run in paragraph.runs[1:]:
                    run.text = ""
            lines.append("PRESENTATION FIXED")

doc.save(path)

# Re-read key rows to a log, avoiding the console encoding limit.
doc = Document(path)
for table in doc.tables:
    for row in table.rows:
        key = row.cells[0].text.strip()
        if key in ("UT-03", "UT-05", "IT-04", "IT-05", "IT-07", "ST-09", "ST-11", "PM-01", "PT-10"):
            lines.append(key + " :: " + " || ".join(cell.text.replace("\n", " ") for cell in row.cells))

for paragraph in doc.paragraphs:
    if paragraph.text.startswith("Figure B10") or paragraph.text.startswith("Figure B11") or paragraph.text.startswith("The results of unit"):
        lines.append("PARA " + paragraph.text[:600])

open(out_log, "w", encoding="utf-8").write("\n\n".join(lines))
print("wrote log", len(lines))
