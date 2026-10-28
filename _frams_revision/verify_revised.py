import docx, re
p = r"C:\Users\HomePC\Downloads\FRAMS_REVISED_MANUSCRIPT.docx"
d = docx.Document(p)
needles = [
    "FRAMS differ", "28 borrower", "three collector", "I am monthly",
    "subject to confirmation", "Khandani et al. (2018)", "Figure L1",
    "observes testing", "eligibility estimates", "FRAMS work",
    "A.1", "October 15", "DTI_after = (E + MP) / I)",
    "one rating per respondent", "criterion-level data",
    "webhook and the paid amount", "verified webhook",
    "4.81", "Presentation of Testing", "APPENDIX B", "APPENDIX L",
    "30 borrowers", "HTML minimum", "Payment Intent",
]
out = r"C:\xampp\htdocs\lending_system\_frams_revision\verify.txt"
with open(out, "w", encoding="utf-8") as f:
    f.write(f"paras {len(d.paragraphs)} tables {len(d.tables)}\n")
    for i, para in enumerate(d.paragraphs):
        t = para.text
        hits = [n for n in needles if n.lower() in t.lower() or n in t]
        if hits:
            f.write(f"\n[{i}] {hits}\n{t[:500]}\n")
    f.write("\n=== CH5 BLOCK ===\n")
    capture = False
    for i, para in enumerate(d.paragraphs):
        t = para.text.strip()
        if t.startswith("CHAPTER V") or t == "Conclusions" or t.startswith("Summary of Findings"):
            capture = True
        if capture:
            f.write(f"[{i}] {t[:400]}\n")
        if capture and t.startswith("APPENDIX"):
            break
        if capture and t.startswith("References") and i > 700:
            f.write(f"[{i}] {t[:200]}\n")
            # continue a bit
    f.write("\n=== TITLE ===\n")
    for i in range(0, 30):
        t = d.paragraphs[i].text.strip()
        if t:
            f.write(f"[{i}] {t}\n")
    f.write("\n=== TABLE0 row5 ===\n")
    t = d.tables[0]
    f.write(t.rows[5].cells[1].text + "\n")
    f.write("\n=== FORMULAS ===\n")
    for i, para in enumerate(d.paragraphs):
        if "Formula 9" in para.text or "DTI_after" in para.text or "I represents" in para.text:
            f.write(f"[{i}] {para.text[:250]}\n")
print("ok")
