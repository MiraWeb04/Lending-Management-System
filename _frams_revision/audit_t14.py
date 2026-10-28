from pathlib import Path
from docx import Document
cands = [
    Path(r"C:\Users\HomePC\Downloads\FRAMS\FRAMS_REVISED_MANUSCRIPT_FINAL.docx"),
    Path(r"C:\Users\HomePC\Downloads\FRAMS_REVISED_MANUSCRIPT_FINAL.docx"),
]
MS = next(p for p in cands if p.exists())
d = Document(str(MS))
print("file", MS)
print("T14 UT-05", d.tables[14].rows[5].cells[1].text if len(d.tables)>14 else "no t14")
# T14 is unit appendix - row 5 might be UT-05
for ti in [7, 14]:
    t = d.tables[ti]
    print("TABLE", ti)
    for row in t.rows:
        if "UT-05" in row.cells[0].text or "UT-05" in row.cells[1].text:
            print(" ", [c.text[:80] for c in row.cells])
for i, p in enumerate(d.paragraphs):
    if "UAT Responses" in p.text or "(1)(2)" in p.text or "corrected criterion" in p.text.lower():
        print(i, p.text[:200])
