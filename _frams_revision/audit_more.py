from pathlib import Path
from docx import Document

cands = [
    Path(r"C:\Users\HomePC\Downloads\FRAMS_REVISED_MANUSCRIPT_FINAL.docx"),
    Path(r"C:\Users\HomePC\Downloads\FRAMS\FRAMS_REVISED_MANUSCRIPT_FINAL.docx"),
]
MS = next(p for p in cands if p.exists())
d = Document(str(MS))
out = Path(r"C:\xampp\htdocs\lending_system\_frams_revision\audit_more.txt")

needles = ("UT-05", "Invalid Document", "Valid JPG", "Background of", "Significance of",
           "PM-07", "borrower role", "Figure B5", "Figure B27", "B26_pm01",
           "How acceptable", "Table 4.", "Figure 5.", "Figure 2", "Figure 3",
           "Appendix A", "APPENDIX A", "APPENDIX C", "APPENDIX I", "APPENDIX J")
with out.open("w", encoding="utf-8") as f:
    f.write(str(MS)+"\n")
    for i, p in enumerate(d.paragraphs):
        t = p.text
        if any(n in t for n in needles):
            f.write(f"[{i}] {t[:280]}\n")
    f.write("\n=== T8 IT ===\n")
    t = d.tables[8]
    for row in t.rows:
        f.write(row.cells[0].text.replace("\n","") + " | " + row.cells[1].text[:70] + " | " + row.cells[4].text[:20] + " | " + row.cells[5].text[:50] + "\n")
    f.write("\n=== T9 ST ids ===\n")
    t = d.tables[9]
    for row in t.rows:
        f.write(row.cells[0].text.replace("\n","") + " | " + row.cells[-1].text[:50] + "\n")
    f.write("\n=== REFS LIST ===\n")
    start = None
    for i, p in enumerate(d.paragraphs):
        if p.text.strip() == "References" and i > 800:
            start = i
            break
    if start:
        for j in range(start, min(start+40, len(d.paragraphs))):
            t = d.paragraphs[j].text.strip()
            if t:
                f.write(f"[{j}] {t[:180]}\n")
            if t.startswith("APPENDIX"):
                break
print("ok")
