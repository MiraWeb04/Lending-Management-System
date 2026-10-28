# -*- coding: utf-8 -*-
from collections import Counter
from pathlib import Path
import re
from openpyxl import load_workbook
from docx import Document

UAT = Path(r"C:\Users\HomePC\Downloads\FRAMS UAT Responses  (1).xlsx")
MS = Path(r"C:\Users\HomePC\Downloads\FRAMS_REVISED_MANUSCRIPT_FINAL.docx")
if not MS.exists():
    MS = Path(r"C:\Users\HomePC\Downloads\FRAMS\FRAMS_REVISED_MANUSCRIPT_FINAL.docx")

out = Path(r"C:\xampp\htdocs\lending_system\_frams_revision\audit_findings.txt")

SCALE = {
    "5 - Strongly Agree": 5,
    "4 - Agree": 4,
    "3 - Neutral": 3,
    "2 - Disagree": 2,
    "1 - Strongly Disagree": 1,
}

wb = load_workbook(UAT, data_only=True)
ws = wb["Form Responses 1"]
headers = [c.value for c in ws[1]]
roles = []
item_vals = []  # 32 x 20
freq = Counter()
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
    if not row or row[2] is None:
        continue
    roles.append(str(row[2]).strip())
    vals = []
    for col in range(3, 23):
        raw = row[col]
        n = SCALE.get(str(raw).strip()) if raw is not None else None
        vals.append(n)
        if n:
            freq[n] += 1
    item_vals.append(vals)

role_count = Counter(roles)
n_resp = len(roles)
assert all(len(v) == 20 and all(x is not None for x in v) for v in item_vals)

criteria = {
    "Functionality": list(range(0, 4)),
    "Usability": list(range(4, 8)),
    "Reliability": list(range(8, 12)),
    "Security": list(range(12, 16)),
    "Performance": list(range(16, 20)),
}

crit_means = {}
crit_freq = {}
for name, idxs in criteria.items():
    scores = []
    cf = Counter()
    for resp in item_vals:
        for i in idxs:
            scores.append(resp[i])
            cf[resp[i]] += 1
    crit_means[name] = sum(scores) / len(scores)
    crit_freq[name] = cf

all_scores = [x for resp in item_vals for x in resp]
overall = sum(all_scores) / len(all_scores)
weighted_score = sum(k * v for k, v in freq.items())

d = Document(str(MS))

needles = [
    r"webhook", r"PM-11", r"49 of 49", r"49 listed", r"twelve PayMongo",
    r"12 PayMongo", r"28 borrowers", r"three collectors", r"30 borrowers",
    r"4\.53", r"4\.56", r"4\.54", r"4\.46", r"4\.52", r"4\.81",
    r"Check with PayMongo", r"Payment Intent", r"live PayMongo",
    r"polling", r"optional webhook", r"verified webhook",
    r"FRAMS differ", r"I am monthly", r"DTI_after = \(E \+ MP\)",
    r"observes testing", r"subject to confirmation", r"Figure L1",
    r"Khandani et al\. \(2018\)", r"PM-01 to PM-12",
    r"11 PayMongo", r"48", r"640", r"341", r"297",
]

hits = {pat: [] for pat in needles}
for i, para in enumerate(d.paragraphs):
    t = para.text
    for pat in needles:
        if re.search(pat, t, re.I):
            hits[pat].append((i, t[:240]))

table_hits = []
for ti, table in enumerate(d.tables):
    blob = "\n".join(c.text for r in table.rows for c in r.cells)
    flags = [pat for pat in needles if re.search(pat, blob, re.I)]
    head = " | ".join(c.text.replace("\n", " ")[:30] for c in table.rows[0].cells)
    table_hits.append((ti, len(table.rows), len(table.columns), flags, head))

# extract objectives and SOP
obj = []
capture = False
for i, para in enumerate(d.paragraphs):
    t = para.text.strip()
    if t in {"Statement of the Problem", "Objectives of the Study", "Scope and Delimitations",
             "CHAPTER V", "Conclusions", "Recommendations", "PayMongo QR Ph Testing"}:
        capture = True
    if capture:
        obj.append(f"[{i}] {t[:300]}")
    if t.startswith("CHAPTER II") or t.startswith("Significance"):
        if i > 50:
            capture = False

with out.open("w", encoding="utf-8") as f:
    f.write(f"MANUSCRIPT {MS}\n")
    f.write(f"PARAS {len(d.paragraphs)} TABLES {len(d.tables)}\n\n")
    f.write("=== UAT RECALC ===\n")
    f.write(f"respondents {n_resp} roles {dict(role_count)}\n")
    f.write(f"item ratings {len(all_scores)}\n")
    f.write(f"freq {dict(freq)} total {sum(freq.values())}\n")
    f.write(f"weighted_score {weighted_score} overall {overall:.6f} rounded {round(overall,2)}\n")
    for name, m in crit_means.items():
        cf = crit_freq[name]
        f.write(f"{name}: mean={m:.6f} round2={round(m,2)} f={dict(cf)} n={sum(cf.values())}\n")
    f.write("\n=== TABLE HEADS ===\n")
    for ti, rows, cols, flags, head in table_hits:
        f.write(f"T{ti} {rows}x{cols} flags={flags}\n  {head}\n")
    f.write("\n=== NEEDLE HITS ===\n")
    for pat, lst in hits.items():
        if lst:
            f.write(f"\n## {pat} ({len(lst)})\n")
            for i, t in lst[:12]:
                f.write(f"  [{i}] {t}\n")
    f.write("\n=== FRONT / OBJECTIVES ===\n")
    for i, para in enumerate(d.paragraphs):
        if i > 80:
            break
        t = para.text.strip()
        if t:
            f.write(f"[{i}] {t[:220]}\n")

print("wrote", out, "overall", overall)
