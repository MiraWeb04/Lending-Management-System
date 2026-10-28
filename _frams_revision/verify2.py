import docx
p = r"C:\Users\HomePC\Downloads\FRAMS_REVISED_MANUSCRIPT.docx"
d = docx.Document(p)
out = r"C:\xampp\htdocs\lending_system\_frams_revision\verify2.txt"
with open(out, "w", encoding="utf-8") as f:
    f.write(f"paras {len(d.paragraphs)} tables {len(d.tables)}\n")
    capture = False
    for i, para in enumerate(d.paragraphs):
        t = para.text.strip()
        if t.startswith("CHAPTER V"):
            capture = True
        if capture:
            f.write(f"[{i}] {t[:300]}\n")
        if capture and t.startswith("Bazarbash"):
            break
    f.write("\n=== NEEDLES ===\n")
    for needle in ["28 borrowers", "three collectors", "I am monthly", "Khandani et al. (2018)", "Figure L1", "subject to confirmation", "A.1", "October 15", "eligibility estimates assigns", "FRAMS differ from", "one rating per respondent", "4.81 for functionality", "installation script that creates every table was not found"]:
        hits = [i for i, para in enumerate(d.paragraphs) if needle in para.text]
        f.write(f"{needle!r}: paras {hits}\n")
        for table in d.tables:
            for ri, row in enumerate(table.rows):
                for ci, cell in enumerate(row.cells):
                    if needle in cell.text:
                        f.write(f"  table cell {ri},{ci}: {cell.text[:120]!r}\n")
    f.write("\n=== APPENDIX L ===\n")
    for i, para in enumerate(d.paragraphs):
        if "APPENDIX L" in para.text or "lending_management" in para.text and "dump" in para.text.lower():
            f.write(f"[{i}] {para.text[:400]}\n")
    f.write("\n=== BOUNDARY ===\n")
    for i, para in enumerate(d.paragraphs):
        if "available payment capacity" in para.text or "HTML minimum" in para.text:
            f.write(f"[{i}] {para.text[:400]}\n")
print("ok")
