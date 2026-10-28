# -*- coding: utf-8 -*-
import re
from pathlib import Path
from datetime import datetime
from docx import Document

FINAL = Path(r"C:\Users\HomePC\Downloads\FRAMS_MANUSCRIPT_FINAL_REVIEWED.docx")
PREV = Path(r"C:\Users\HomePC\Downloads\FRAMS_MANUSCRIPT_AUDIT_CORRECTED.docx")
OUT = Path(r"C:\xampp\htdocs\lending_system\_frams_revision\final_reopen_verify.txt")

d = Document(str(FINAL))
prev = Document(str(PREV))
lines = []
lines.append(f"FINAL {FINAL} size={FINAL.stat().st_size} mtime={datetime.fromtimestamp(FINAL.stat().st_mtime)}")
lines.append(f"PREV  {PREV} size={PREV.stat().st_size} mtime={datetime.fromtimestamp(PREV.stat().st_mtime)}")
lines.append(f"FINAL paras={len(d.paragraphs)} tables={len(d.tables)}")
lines.append(f"PREV  paras={len(prev.paragraphs)} tables={len(prev.tables)}")
prev_body = "\n".join(p.text for p in prev.paragraphs)
lines.append(f"PREV still lacks Santos, A. R.={('Santos, A. R. (2023)' not in prev_body)}")
lines.append(f"PREV still has broken T21 header Strongly Agree twice={prev.tables[21].rows[0].cells[0].text.strip()=='Strongly Agree'}")

body = "\n".join(p.text for p in d.paragraphs)
# cites
patterns = re.findall(
    r"([A-Z][A-Za-z'’\-]+(?:\s+(?:and|&)\s+[A-Z][A-Za-z'’\-]+)?(?:\s+et al\.)?)\s*\((\d{4}|n\.d\.)\)",
    body.split("References")[0] if "References" in body else body,
)
patterns2 = re.findall(
    r"\(([A-Z][A-Za-z'’\-]+(?:\s+(?:&|and)\s+[A-Z][A-Za-z'’\-]+)?(?:\s+et al\.)?),\s*(\d{4}|n\.d\.)\)",
    body.split("References")[0] if "References" in body else body,
)
uniq = sorted(set((a.strip(), y) for a, y in patterns + patterns2))
lines.append("\n=== IN-TEXT ===")
for a, y in uniq:
    lines.append(f"  {a} ({y})")

# ref authors
in_refs = False
refs = []
for i, p in enumerate(d.paragraphs):
    t = p.text.strip()
    if t == "References" and i > 800:
        in_refs = True
        continue
    if in_refs:
        if t.startswith("APPENDIX"):
            break
        if t and t != "The following references are presented in APA 7th edition format.":
            refs.append(t)
            lines.append(f"REF [{i}] {t[:220]}")

def ref_covers(author, year):
    blob = "\n".join(refs)
    if author == "Santos":
        return "Santos, A. R. (2023)" in blob
    if author == "BlueOrchard":
        return "BlueOrchard. (2025" in blob
    if author == "LendStack":
        return "LendStack. (n.d.)" in blob
    if author.startswith("Bazarbash"):
        return "Bazarbash" in blob and "2020" in blob
    if author.startswith("Campos"):
        return "Campos" in blob
    if author == "Davis":
        return "Davis, F. D. (1989)" in blob
    if author.startswith("DeLone"):
        return "DeLone" in blob and "2003" in blob
    if author.startswith("Estrellado"):
        return "Estrellado" in blob
    if author.startswith("Juventud"):
        return "Juventud" in blob
    if author.startswith("Khandani"):
        return "Khandani" in blob
    if author == "Mina":
        return "Mina, J." in blob
    if author.startswith("Penaso"):
        return "Penaso" in blob
    if author.startswith("Ramirez"):
        return "Ramirez" in blob
    if author == "PayMongo":
        return "PayMongo. (n.d.)" in blob
    return False

lines.append("\n=== MATCH ===")
missing = []
for a, y in uniq:
    key = a.replace(" and ", " ").replace("& ", "")
    covered = ref_covers(a.split()[0] if a.split()[0] not in ("Chen",) else a, y)
    # simpler
    last = a.split()[0]
    hit = any(last.replace(",", "") in r for r in refs)
    if last in ("Chen", "Frost", "Ghosh", "Juszczak", "Lopez", "Rahman", "Kou", "Batista"):
        hit = False
    if last == "ISO":
        hit = True
    lines.append(f"  {a} ({y}) -> {'MATCH' if hit else 'MISSING'}")
    if not hit:
        missing.append(f"{a} ({y})")

lines.append("\nStill missing: " + "; ".join(missing))

# invariants
for needle in [
    "30 borrowers, one administrator, and one collector",
    "640 individual",
    "overall weighted mean was 4.53",
    "4.53",
]:
    lines.append(f"INV {needle!r}: {body.lower().count(needle.lower()) if False else body.count(needle) if needle!='4.53' else len(re.findall(r'4\\.53', body))}")

t12 = d.tables[12]
lines.append("T12 total: " + " | ".join(c.text.replace("\n"," ")[:90] for c in t12.rows[-1].cells))
t13 = d.tables[13]
lines.append("T13 overall: " + " | ".join(c.text.strip() for c in t13.rows[-1].cells))
t21 = d.tables[21]
lines.append(f"T21 {len(t21.rows)}x{len(t21.columns)} header={ [c.text.strip() for c in t21.rows[0].cells] }")
nums = [row.cells[0].text.strip() for row in t21.rows if row.cells[0].text.strip().isdigit()]
lines.append(f"T21 item nos={nums}")

# caption locations
for i, p in enumerate(d.paragraphs):
    t = p.text.strip()
    if t.startswith("Table I1") or t.startswith("Figure B32") or "LendStack. (n.d.)" in t:
        lines.append(f"LOC [{i}] {t[:180]}")

# official 48 and webhook exclusion
lines.append(f"PM-11 excluded note present={('excluded from the official PayMongo' in body) or ('excluded from official PayMongo' in body)}")
lines.append(f"Check with PayMongo mentioned={body.count('Check with PayMongo')}")
lines.append(f"live PayMongo leftover overclaim check: 'live PayMongo charge' count={body.count('live PayMongo')}")

OUT.write_text("\n".join(lines), encoding="utf-8")
print("\n".join(lines[-40:]))
print("wrote", OUT)
