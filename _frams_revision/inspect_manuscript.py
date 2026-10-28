import docx

p = r"C:\Users\HomePC\Downloads\FRAMS_3RDREVISIONS.docx"
d = docx.Document(p)
out = r"C:\xampp\htdocs\lending_system\_frams_revision\struct.txt"

def has_drawing(para):
    xml = para._p.xml
    return ("w:drawing" in xml) or ("w:pict" in xml)

with open(out, "w", encoding="utf-8") as f:
    f.write("=== METHODOLOGY 165-185 ===\n")
    for i in range(165, 186):
        t = d.paragraphs[i].text.strip().replace("\n", " ")
        f.write(f"[{i}] {t[:400]}\n")

    f.write("\n=== APPENDIX MARKERS ===\n")
    for i, para in enumerate(d.paragraphs):
        if i < 840:
            continue
        t = para.text.strip().replace("\n", " ")
        if not t:
            continue
        if len(t) < 100 or t.upper().startswith("APPENDIX") or t.startswith("CHAPTER"):
            f.write(f"[{i}|d={int(has_drawing(para))}] {t[:160]}\n")

    f.write("\n=== TABLES ===\n")
    for ti, t in enumerate(d.tables):
        rows = len(t.rows)
        cols = len(t.columns)
        first = " | ".join(c.text.strip().replace("\n", " ")[:36] for c in t.rows[0].cells)
        blob = "\n".join(c.text for row in t.rows for c in row.cells)
        flags = []
        for key in ["Khandani", "2018", "28 borrower", "three collector", "30 borrower", "UT-0", "1,000", "I am monthly", "DTI_after", "editorial"]:
            if key.lower() in blob.lower():
                flags.append(key)
        f.write(f"T{ti} {rows}x{cols} flags={flags}\n  {first[:200]}\n")

    f.write("\n=== TARGET DRAW ===\n")
    for i in [4, 239, 256, 258, 286, 302, 303, 350, 397, 399, 403, 458]:
        para = d.paragraphs[i]
        f.write(f"[{i}] draw={int(has_drawing(para))} runs={len(para.runs)} {para.text[:160]!r}\n")

    f.write("\n=== FORMULAS ===\n")
    for i in range(548, 600):
        t = d.paragraphs[i].text.strip()
        if t:
            f.write(f"[{i}] {t[:500]}\n")

    f.write("\n=== CH5 ===\n")
    for i in range(800, 870):
        t = d.paragraphs[i].text.strip().replace("\n", " ")
        if t:
            f.write(f"[{i}] {t[:420]}\n")

print("ok")
