from pathlib import Path
from docx import Document
import re

# prefer the copy that exists
cands = [
    Path(r"C:\Users\HomePC\Downloads\FRAMS_REVISED_MANUSCRIPT_FINAL.docx"),
    Path(r"C:\Users\HomePC\Downloads\FRAMS\FRAMS_REVISED_MANUSCRIPT_FINAL.docx"),
]
MS = next(p for p in cands if p.exists())
print("USING", MS)
d = Document(str(MS))
out = Path(r"C:\xampp\htdocs\lending_system\_frams_revision\audit_tables.txt")

def dump_table(f, ti):
    t = d.tables[ti]
    f.write(f"\n===== T{ti} {len(t.rows)}x{len(t.columns)} =====\n")
    for ri, row in enumerate(t.rows):
        f.write(f"R{ri}: {' || '.join(c.text.replace(chr(10),' / ')[:200] for c in row.cells)}\n")

with out.open("w", encoding="utf-8") as f:
    f.write(str(MS) + "\n")
    for ti in [1, 7, 10, 12, 13, 19, 20, 22, 23]:
        dump_table(f, ti)

    f.write("\n=== T21 first 3 rows ===\n")
    t = d.tables[21]
    f.write(f"{len(t.rows)}x{len(t.columns)}\n")
    for ri in range(min(3, len(t.rows))):
        f.write(f"R{ri}: {[c.text[:40] for c in t.rows[ri].cells]}\n")

    f.write("\n=== OBJECTIVES / SOP / CONCLUSIONS ===\n")
    keys = ("Statement of the Problem", "Objectives of the Study", "Specifically",
            "Scope and Delimitations", "CHAPTER V", "Summary of Findings",
            "Conclusions", "Recommendations", "Background of")
    capture = False
    for i, p in enumerate(d.paragraphs):
        t = p.text.strip()
        if any(t.startswith(k) or t == k for k in keys) or t.startswith("How can FRAMS") or t.startswith("Develop ") or t.startswith("Implement ") or t.startswith("Evaluate ") or t.startswith("In relation to"):
            f.write(f"[{i}] {t[:350]}\n")

    f.write("\n=== T7 unit IDs ===\n")
    t = d.tables[7]
    for row in t.rows:
        f.write(row.cells[0].text.replace("\n","") + " | " + row.cells[1].text[:80] + " | " + row.cells[4].text[:40] + " | " + row.cells[5].text[:60] + "\n")

    f.write("\n=== CITATIONS vs REFS ===\n")
    body = "\n".join(p.text for p in d.paragraphs)
    cites = sorted(set(re.findall(r"([A-Z][A-Za-z'’\-]+(?:\s+(?:and|&)\s+[A-Z][A-Za-z'’\-]+)?(?:\s+et al\.)?)\s*\((\d{4}|n\.d\.)\)", body)))
    f.write("sample cites count %d\n" % len(cites))
    for c in cites[:80]:
        f.write(f"  {c[0]} ({c[1]})\n")

    f.write("\n=== REFS START ===\n")
    for i, p in enumerate(d.paragraphs):
        if p.text.strip() in {"References", "REFERENCES"} or p.text.strip().startswith("Bazarbash"):
            for j in range(i, min(i+25, len(d.paragraphs))):
                if d.paragraphs[j].text.strip():
                    f.write(f"[{j}] {d.paragraphs[j].text[:200]}\n")
            break

print("ok")
