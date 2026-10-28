import docx
from docx.oxml.ns import qn
from docx.shared import Pt
from pathlib import Path
import shutil

p = Path(r"C:\Users\HomePC\Downloads\FRAMS_REVISED_MANUSCRIPT.docx")
d = docx.Document(str(p))
para = d.paragraphs[625]
print(repr(para.text))
print("---")
# also nearby
for i in range(620, 630):
    print(i, d.paragraphs[i].text[:250])
