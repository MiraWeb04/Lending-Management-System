# -*- coding: utf-8 -*-
import re
from pathlib import Path
from docx import Document

p = Path(r"C:\Users\HomePC\Downloads\FRAMS_REVISED_MANUSCRIPT.docx")
d = Document(str(p))
out = Path(r"C:\xampp\htdocs\lending_system\_frams_revision\webhook_hits.txt")

needles = re.compile(
    r"webhook|WebHook|WEBHOOK|payment\.paid|evt_|signature.?valid|invalid_signature|"
    r"PM-09|PM-10|PM-11|PM-12|Table 4\.7|Table 4\.9|49 of 49|49 passed|"
    r"twelve PayMongo|12 PayMongo|PM-01",
    re.I,
)

with out.open("w", encoding="utf-8") as f:
    f.write(f"PARAS {len(d.paragraphs)} TABLES {len(d.tables)}\n\n")
    f.write("=== PARAGRAPHS ===\n")
    for i, para in enumerate(d.paragraphs):
        t = para.text
        if needles.search(t):
            f.write(f"\n[{i}|{para.style.name if para.style else ''}]\n{t}\n")

    f.write("\n\n=== TABLES ===\n")
    for ti, table in enumerate(d.tables):
        blob = "\n".join(c.text for row in table.rows for c in row.cells)
        if needles.search(blob):
            f.write(f"\n----- TABLE {ti} {len(table.rows)}x{len(table.columns)} -----\n")
            for ri, row in enumerate(table.rows):
                cells = [c.text.replace("\n", " / ")[:160] for c in row.cells]
                f.write(f"R{ri}: {' || '.join(cells)}\n")

print("wrote", out)
