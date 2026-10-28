from pathlib import Path
from docx import Document

d = Document(r"C:\Users\HomePC\Downloads\FRAMS_REVISED_MANUSCRIPT.docx")
out = Path(r"C:\xampp\htdocs\lending_system\_frams_revision\dump_more.txt")
with out.open("w", encoding="utf-8") as f:
    for i in [32, 69, 110, 229, 232, 239, 242, 271, 279, 282, 458, 459, 492, 493, 505, 644, 698, 699, 700, 716, 838, 839, 845, 849, 855]:
        if i < len(d.paragraphs):
            f.write(f"\n[{i}]\n{d.paragraphs[i].text}\n")
    f.write("\n=== T6 payment row ===\n")
    t = d.tables[6]
    for ri, row in enumerate(t.rows):
        blob = " | ".join(c.text[:80] for c in row.cells)
        if "Pay" in blob or "pay" in blob or "QR" in blob:
            f.write(f"R{ri}: {blob}\n")
    f.write("\n=== T12 full ===\n")
    t = d.tables[12]
    for ri, row in enumerate(t.rows):
        f.write(f"R{ri}: {' || '.join(c.text.replace(chr(10),' / ')[:200] for c in row.cells)}\n")
    f.write("\n=== T20 full ===\n")
    t = d.tables[20]
    for ri, row in enumerate(t.rows):
        f.write(f"R{ri}: {' || '.join(c.text.replace(chr(10),' / ')[:200] for c in row.cells)}\n")
    f.write("\n=== T10 PM-09 to PM-12 full ===\n")
    t = d.tables[10]
    for ri in range(9, 13):
        f.write(f"R{ri}:\n")
        for ci, cell in enumerate(t.rows[ri].cells):
            f.write(f"  C{ci}: {cell.text}\n")
print("ok")
