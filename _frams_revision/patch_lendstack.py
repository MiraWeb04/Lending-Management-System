# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_ALIGN_PARAGRAPH

DEST = Path(r"C:\Users\HomePC\Downloads\FRAMS_MANUSCRIPT_FINAL_REVIEWED.docx")
CHANGES = Path(r"C:\Users\HomePC\Downloads\FRAMS_FINAL_REVIEW_CHANGE_SUMMARY.docx")

def set_run_font(run, name="Times New Roman", size=12, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic

def insert_paragraph_after(paragraph):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)

doc = Document(str(DEST))
body = "\n".join(p.text for p in doc.paragraphs)
print("HAS LendStack. (n.d.) REF", "LendStack. (n.d.)" in body)
print("HAS LendStack (n.d.)", "LendStack (n.d.)" in body)
print("HAS LendStack, n.d.", "LendStack, n.d." in body)
for i, p in enumerate(doc.paragraphs):
    if "LendStack" in p.text:
        print(f"CITE [{i}] {p.text[:250]}")

if "LendStack. (n.d.)" not in body:
    target = None
    for p in doc.paragraphs:
        if p.text.strip().startswith("Khandani, A."):
            target = p
            break
    if target is None:
        raise SystemExit("Khandani not found")
    para = insert_paragraph_after(target)
    parts = [
        ("LendStack. (n.d.). ", False, False),
        ("Building the future of lending infrastructure", False, True),
        (". https://www.lendstack.in/", False, False),
    ]
    for text, bold, italic in parts:
        run = para.add_run(text)
        set_run_font(run, size=12, bold=bold, italic=italic)
    fmt = para.paragraph_format
    fmt.left_indent = Inches(0.5)
    fmt.first_line_indent = Inches(-0.5)
    fmt.space_after = Pt(6)
    fmt.line_spacing = 2.0
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.save(str(DEST))
    print("INSERTED LendStack reference after Khandani")
else:
    print("LendStack ref already present")

# Table I1 caption
d2 = Document(str(DEST))
found_caption = any("Table I1" in p.text for p in d2.paragraphs)
print("Table I1 caption", found_caption)
print("PayMongo, n.d. in-text", "(PayMongo, n.d.)" in "\n".join(p.text for p in d2.paragraphs))

# change summary applied bullets
if CHANGES.exists():
    c = Document(str(CHANGES))
    print("\n=== CHANGE SUMMARY ===")
    for p in c.paragraphs:
        t = p.text.strip()
        if t:
            print(t[:300])
