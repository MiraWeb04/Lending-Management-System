from docx import Document
from docx.oxml.ns import qn

doc = Document(r"testing_evidence\_harness\FRAMS_REVISED_DRAFT_BEFORE_TEST_ALIGN.docx")
lines = []
for i in range(1040, 1160):
    p = doc.paragraphs[i]
    has_drawing = bool(p._p.xpath(".//*[local-name()='drawing']"))
    text = p.text.replace("\n", " ")[:160]
    if text or has_drawing:
        lines.append(f"P{i} draw={int(has_drawing)} {text}")
lines.append("--- second copy ---")
for i in range(1300, 1420):
    p = doc.paragraphs[i]
    has_drawing = bool(p._p.xpath(".//*[local-name()='drawing']"))
    text = p.text.replace("\n", " ")[:160]
    if text or has_drawing:
        lines.append(f"P{i} draw={int(has_drawing)} {text}")
open(r"_frams_revision\dump_figures.txt", "w", encoding="utf-8").write("\n".join(lines))
print("lines", len(lines))
