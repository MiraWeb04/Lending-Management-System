# -*- coding: utf-8 -*-
import re
from pathlib import Path
from docx import Document

MS = Path(r"C:\Users\HomePC\Downloads\FRAMS_MANUSCRIPT_AUDIT_CORRECTED.docx")
d = Document(str(MS))
out = Path(r"C:\xampp\htdocs\lending_system\_frams_revision\cites.txt")

body_paras = []
refs = []
in_refs = False
for i, p in enumerate(d.paragraphs):
    t = p.text.strip()
    if t == "References" and i > 800:
        in_refs = True
    if in_refs:
        if t.startswith("APPENDIX"):
            break
        if t:
            refs.append((i, t))
    else:
        body_paras.append((i, t))

body = "\n".join(t for _, t in body_paras)
# citations
patterns = re.findall(
    r"([A-Z][A-Za-z'’\-]+(?:\s+(?:and|&)\s+[A-Z][A-Za-z'’\-]+)?(?:\s+et al\.)?)\s*\((\d{4}|n\.d\.)\)",
    body,
)
# also (Author, year)
patterns2 = re.findall(
    r"\(([A-Z][A-Za-z'’\-]+(?:\s+(?:&|and)\s+[A-Z][A-Za-z'’\-]+)?(?:\s+et al\.)?),\s*(\d{4}|n\.d\.)\)",
    body,
)
allc = patterns + patterns2
uniq = sorted(set((a.strip(), y) for a, y in allc))

with out.open("w", encoding="utf-8") as f:
    f.write("=== BODY CITES ===\n")
    for a, y in uniq:
        f.write(f"{a}\t{y}\n")
    f.write("\n=== REF LIST ===\n")
    for i, t in refs:
        f.write(f"[{i}] {t}\n")
    f.write("\n=== APPENDIX I nearby ===\n")
    for i, p in enumerate(d.paragraphs):
        if "APPENDIX I" in p.text or (i >= 1240 and i <= 1285):
            f.write(f"[{i}|{p.style.name if p.style else ''}] {p.text[:200]}\n")
    f.write(f"\nT21 {len(d.tables[21].rows)}x{len(d.tables[21].columns)}\n")
    t = d.tables[21]
    for ri, row in enumerate(t.rows[:6]):
        f.write(f"R{ri}: {[c.text[:50].replace(chr(10),' | ') for c in row.cells]}\n")
print("cites", len(uniq), "refs", len(refs))
