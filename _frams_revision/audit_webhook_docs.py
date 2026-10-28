import re
from docx import Document

paths = [
    r"testing_evidence\_harness\FRAMS_REVISED_DRAFT_BEFORE_TEST_ALIGN.docx",
    r"lending_system\documentation\Test_Cases_PayMongo_Penetration_and_System_Testing.docx",
    r"testing_evidence\FRAMS_QRPh_Figures_B26_B37_updated.docx",
    r"testing_evidence\FINAL_TESTING_REPORT.docx",
    r"lending_system\documentation\FRAMS_References_Questionnaires_and_Forms.docx",
]

for path in paths:
    print("=" * 80)
    print(path)
    doc = Document(path)
    print("paragraphs", len(doc.paragraphs), "tables", len(doc.tables))
    for i, p in enumerate(doc.paragraphs):
        if re.search(r"webhook|PAYMONGO_WEBHOOK|PM-09|PM-10|PM-11|Payment Intent|polling", p.text, re.I):
            print(f"P{i}: {p.text[:500].replace(chr(10), ' ')}")
    for ti, table in enumerate(doc.tables):
        blob = "\n".join(" | ".join(c.text.strip() for c in row.cells) for row in table.rows)
        if re.search(r"webhook|PayMongo|PM-", blob, re.I):
            print(f"--- TABLE {ti} rows={len(table.rows)} cols={len(table.columns)} ---")
            for ri, row in enumerate(table.rows):
                cells = [c.text.strip().replace("\n", " ")[:180] for c in row.cells]
                line = " || ".join(cells)
                if re.search(r"webhook|PM-|PayMongo|Total|Unit|System", line, re.I):
                    print(f"R{ri}: {line[:500]}")
    print()
