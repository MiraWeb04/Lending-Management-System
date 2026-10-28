# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document

MS = Path(r"C:\Users\HomePC\Downloads\FRAMS_MANUSCRIPT_FINAL_REVIEWED.docx")
d = Document(str(MS))
out = Path(r"C:\xampp\htdocs\lending_system\_frams_revision\ch2_final.txt")
lines = []
for i, p in enumerate(d.paragraphs):
    if 113 <= i <= 175 or (i >= 868 and i <= 895):
        t = p.text.strip()
        if t:
            lines.append(f"[{i}] {t}")
# table 1
lines.append("\n=== T0 ===")
t = d.tables[0]
for ri, row in enumerate(t.rows):
    cells = [c.text.replace("\n"," ").strip()[:120] for c in row.cells]
    lines.append(f"R{ri}: {cells}")
out.write_text("\n".join(lines), encoding="utf-8")
print("ok", len(lines))
