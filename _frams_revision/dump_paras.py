from docx import Document

doc = Document(r"testing_evidence\_harness\FRAMS_REVISED_DRAFT_BEFORE_TEST_ALIGN.docx")
wanted = [36, 67, 217, 227, 279, 471, 506, 507, 515, 666, 720, 1002, 1259]
out = []
for i in wanted:
    out.append(f"\n===== P{i} =====\n")
    out.append(doc.paragraphs[i].text)
print("TABLE 11")
t = doc.tables[11]
for ri, row in enumerate(t.rows):
    cells = [c.text.replace("\n", " | ") for c in row.cells]
    out.append(f"T11R{ri}: " + " || ".join(cells))
print("TABLE 19")
t = doc.tables[19]
for ri, row in enumerate(t.rows):
    cells = [c.text.replace("\n", " | ") for c in row.cells]
    out.append(f"T19R{ri}: " + " || ".join(cells))
# find all 49 and twelve
for i, p in enumerate(doc.paragraphs):
    if "49" in p.text and ("test" in p.text.lower() or "passed" in p.text.lower()):
        out.append(f"\n===== COUNT P{i} =====\n{p.text[:800]}")
    if "twelve PayMongo" in p.text or "12 PayMongo" in p.text or "PM-01 to PM-12" in p.text:
        out.append(f"\n===== RANGE P{i} =====\n{p.text[:800]}")
open(r"_frams_revision\dump_paras.txt", "w", encoding="utf-8").write("\n".join(out))
print("wrote", len(out))
