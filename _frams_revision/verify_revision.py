import re
import zipfile
from docx import Document

path = r"manuscript_revision\FRAMS_Manuscript_REVISED.docx"
doc = Document(path)
print("paragraphs", len(doc.paragraphs), "tables", len(doc.tables))
needles = ["webhook", "PM-10", "PM-11", "PM-12", "B47", "B48", "B49", "49 functional", "twelve PayMongo", "12/12 cases passed"]
blob_parts = []
for i, p in enumerate(doc.paragraphs):
    blob_parts.append(p.text)
    for n in needles:
        if n.lower() in p.text.lower():
            print(f"PARA {i} has {n}: {p.text[:180]}")
for ti, table in enumerate(doc.tables):
    ids = [row.cells[0].text.strip() for row in table.rows]
    if any(x.startswith("PM-") for x in ids):
        print("TABLE", ti, "ids", ids)
    label = table.rows[0].cells[0].text.strip()
    if label in ("Testing Level", "Testing level"):
        for row in table.rows:
            print("SUMMARY", ti, " | ".join(c.text.replace("\n", " ")[:80] for c in row.cells))
    for row in table.rows:
        joined = " | ".join(c.text for c in row.cells)
        for n in needles:
            if n.lower() in joined.lower():
                print(f"CELL T{ti} has {n}: {joined[:180]}")

# figure sequence sample
text = "\n".join(blob_parts)
for token in ["B26", "B33", "B34", "B35", "B46", "B37", "B38"]:
    print(token, text.count(token))
print("PM-09 count", text.count("PM-09"))
print("succeeded-intent" in text or "succeeded" in text)
