# -*- coding: utf-8 -*-
from pathlib import Path
import shutil
import re
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm

FINAL = Path(r"C:\Users\HomePC\Downloads\FRAMS_REVISED_MANUSCRIPT_FINAL.docx")
REPORT = Path(r"C:\Users\HomePC\Downloads\FRAMS_PAYMONGO_SCOPE_REPORT.docx")
WORKSPACE = Path(r"C:\xampp\htdocs\lending_system\_frams_revision")


def set_run_font(run, name="Times New Roman", size=12, bold=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def replace_text(paragraph, new_text):
    drawing_runs = []
    for run in paragraph.runs:
        xml = run._r.xml
        if "w:drawing" in xml or "w:pict" in xml:
            drawing_runs.append(run._r)
    for run in list(paragraph.runs):
        run._r.getparent().remove(run._r)
    for el in drawing_runs:
        paragraph._p.append(el)
    run = paragraph.add_run(new_text)
    set_run_font(run)


def set_cell(cell, text):
    if not cell.paragraphs:
        return
    replace_text(cell.paragraphs[0], text)
    for extra in cell.paragraphs[1:]:
        extra.text = ""


doc = Document(str(FINAL))

# Fix Table 20 total "Not yet tested" leftover remarks.
t20 = doc.tables[20]
set_cell(t20.rows[6].cells[4], "0")
set_cell(t20.rows[6].cells[5], "0")

# Soften Table 12 PayMongo remarks to avoid the word webhook.
t12 = doc.tables[12]
set_cell(
    t12.rows[4].cells[4],
    "11/11 official cases passed in the isolated test environment. PM-01/02 verified safe rejection with an invalid test key. PM-09/10 document receipt and payment-list display for a stored record. The former local-endpoint signature case is excluded. Polling and staff Check with PayMongo are implemented but not separately passed. See Figures B26–B35 and B37.",
)

# Paragraph wording: keep schema names, avoid extra webhook testing language.
replacements = [
    (
        "Webhook signature validation (formerly PM-11, Figure B36) is excluded from the official count because webhook processing is outside the selected confirmation workflow. No live PayMongo charge is claimed.",
        "The former local-endpoint signature case (Figure B36) is excluded from the official count because it is outside the selected Payment Intent polling and staff Check with PayMongo workflow. No live PayMongo charge is claimed.",
    ),
    (
        "Formerly numbered PM-11 (Figure B36) is excluded from Table 4.7 because it tested local webhook signatures, which are outside the selected confirmation workflow.",
        "Formerly numbered PM-11 (Figure B36) is excluded from Table 4.7 because it tested local endpoint signatures, which are outside the selected confirmation workflow.",
    ),
    (
        "Figure B36 (B36_pm11_webhook_responses.png)",
        "Figure B36 (excluded from official PayMongo scope)",
    ),
    (
        "Figure B36 (B36_pm11_webhook_responses.png) ",
        "Figure B36 (excluded from official PayMongo scope)",
    ),
]

for para in doc.paragraphs:
    t = para.text
    new = t
    for old, repl in replacements:
        if old in new:
            new = new.replace(old, repl)
    if new != t:
        replace_text(para, new)

doc.save(str(FINAL))

check = Document(str(FINAL))
rx = re.compile(r"webhook|invalid_signature|evt_local", re.I)
hits = []
for i, para in enumerate(check.paragraphs):
    if rx.search(para.text):
        hits.append(f"P{i}: {para.text.strip()[:280]}")
for ti, table in enumerate(check.tables):
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            if rx.search(cell.text):
                hits.append(f"T{ti}R{ri}C{ci}: {cell.text.strip()[:280]}")

# count leftover 49 in testing context
count49 = []
for i, para in enumerate(check.paragraphs):
    if "49" in para.text and ("test" in para.text.lower() or "PayMongo" in para.text):
        count49.append(f"P{i}: {para.text[:200]}")
for ti, table in enumerate(check.tables):
    blob = "\n".join(c.text for r in table.rows for c in r.cells)
    if "49 listed" in blob or "49 of 49" in blob:
        count49.append(f"T{ti}: contains old 49 testing total")

out = Path(r"C:\xampp\htdocs\lending_system\_frams_revision\final_webhook_search.txt")
out.write_text("HITS\n" + "\n\n".join(hits) + "\n\nCOUNT49\n" + "\n".join(count49), encoding="utf-8")
print("hits", len(hits))
print("count49", count49)
shutil.copy2(FINAL, WORKSPACE / FINAL.name)
