# -*- coding: utf-8 -*-
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph
from pathlib import Path
import shutil

REVISED = Path(r"C:\Users\HomePC\Downloads\FRAMS_REVISED_MANUSCRIPT.docx")
WORKSPACE = Path(r"C:\xampp\htdocs\lending_system\_frams_revision") / REVISED.name


def set_run_font(run, name="Times New Roman", size=12, bold=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def replace_text(paragraph, new_text):
    drawing_runs = []
    for run in paragraph.runs:
        xml = run._r.xml
        if "w:drawing" in xml or "w:pict" in xml:
            drawing_runs.append(run._r)
    for run in list(paragraph.runs):
        run._r.getparent().remove(run._r)
    p = paragraph._p
    for el in drawing_runs:
        p.append(el)
    run = paragraph.add_run(new_text)
    set_run_font(run)


def insert_after(paragraph, text) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    run = new_para.add_run(text)
    set_run_font(run)
    new_para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    new_para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    new_para.paragraph_format.first_line_indent = Cm(1.27)
    return new_para


def delete_para(paragraph):
    el = paragraph._element
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)


CONCLUSIONS = [
    "In relation to the first specific objective, FRAMS provides a borrower management module that supports registration, administrator approval, client records, online loan applications, and supporting document submission and review.",
    "In relation to the second specific objective, FRAMS computes a rule-based eligibility estimate from declared income, existing obligations, requested amount, loan term, and configured limits. The estimate is a decision-support result and does not automatically approve or reject a loan.",
    "In relation to the third specific objective, FRAMS supports interest computation, installment schedules, loan agreement acceptance, collector assignment, loan release, payment recording, penalty computation, and repayment monitoring. PayMongo QR Ph repayments are recorded after Payment Intent confirmation, not merely because a QR image was displayed.",
    "In relation to the fourth specific objective, FRAMS sends in-application notifications for relevant application decisions, agreement updates, loan releases, and payment transactions, consistent with the notification functions available in the developed system.",
    "In relation to the fifth specific objective, FRAMS generates financial reports that support monitoring of loan releases, collections, outstanding balances, overdue accounts, and related lending transactions.",
    "In relation to the sixth specific objective, 32 intended users evaluated FRAMS. The supplied workbook contains 30 borrowers, one administrator, and one collector, each of whom answered 20 questionnaire items. The overall weighted mean was 4.53, interpreted as Strongly Agree. These ratings are user perceptions of acceptability and are not independent proof of security, loan accuracy, or measured performance.",
    "The documented technical tests support the specific cases examined in their documented environments. Passing those cases does not establish that the entire system is free of defects. Local PayMongo simulations do not independently demonstrate a successful external PayMongo charge or a provider-delivered webhook.",
    "The conclusions of this study are therefore limited to the system functions implemented, the documented technical tests, and the recorded user evaluation responses.",
]


def main():
    doc = Document(str(REVISED))

    # Rebuild conclusions between Conclusions and Recommendations.
    concl_heading = None
    rec_heading = None
    for para in doc.paragraphs:
        t = para.text.strip()
        if t == "Conclusions":
            concl_heading = para
        elif t.startswith("Recommendations") and concl_heading is not None:
            rec_heading = para
            break
    if concl_heading is None or rec_heading is None:
        raise SystemExit("Could not locate Conclusions/Recommendations headings")

    # Delete existing conclusion body
    body = doc.element.body
    children = list(body.iterchildren())
    start = children.index(concl_heading._element)
    stop = children.index(rec_heading._element)
    for el in children[start + 1:stop]:
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)

    anchor = concl_heading
    for text in reversed(CONCLUSIONS):
        # insert_after puts immediately after anchor, so reverse to keep order
        pass
    # Insert in order: each new para after previous
    prev = concl_heading
    for text in CONCLUSIONS:
        prev = insert_after(prev, text)

    # Remove duplicate References heading that is exactly "References" after the APA note
    refs = [p for p in doc.paragraphs if p.text.strip() == "References"]
    if len(refs) >= 2:
        # keep the first, delete later exact duplicates before bibliography entries
        for extra in refs[1:]:
            delete_para(extra)

    doc.save(str(REVISED))
    shutil.copy2(REVISED, WORKSPACE)
    check = Document(str(REVISED))
    print("paras", len(check.paragraphs), "tables", len(check.tables))
    capture = False
    for p in check.paragraphs:
        t = p.text.strip()
        if t == "Conclusions" or t.startswith("CHAPTER V"):
            capture = True
        if capture:
            print(t[:180])
        if t.startswith("Bazarbash"):
            break


if __name__ == "__main__":
    main()
