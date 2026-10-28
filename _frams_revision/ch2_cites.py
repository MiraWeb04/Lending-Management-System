from pathlib import Path
from docx import Document
MS = Path(r"C:\Users\HomePC\Downloads\FRAMS_MANUSCRIPT_AUDIT_CORRECTED.docx")
d = Document(str(MS))
out = Path(r"C:\xampp\htdocs\lending_system\_frams_revision\ch2.txt")
with out.open("w", encoding="utf-8") as f:
    for i, p in enumerate(d.paragraphs):
        if 113 <= i <= 175:
            t = p.text.strip()
            if t:
                f.write(f"[{i}] {t}\n\n")
        if i == 0:
            f.write("start\n")
print("ok", len(d.paragraphs))
