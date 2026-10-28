from docx import Document
from pathlib import Path

files = [
    Path(r"C:\Users\HomePC\Downloads\FRAMS_3RDREVISIONS.docx"),
    Path(r"C:\Users\HomePC\Downloads\FRAMS_REVISED_MANUSCRIPT.docx"),
    Path(r"C:\Users\HomePC\Downloads\FRAMS_REVISION_REPORT.docx"),
    Path(r"C:\Users\HomePC\Downloads\FRAMS_UNRESOLVED_ISSUES.docx"),
    Path(r"C:\Users\HomePC\Downloads\FRAMS_COMPARISON_REPORT.docx"),
]
out = Path(r"C:\xampp\htdocs\lending_system\_frams_revision\final_open_check.txt")
lines = []
for f in files:
    d = Document(str(f))
    first = next((p.text.strip() for p in d.paragraphs if p.text.strip()), "")
    lines.append(f"{f.name}: paras={len(d.paragraphs)} tables={len(d.tables)} first={first[:80]!r}")
orig = Document(str(files[0]))
rev = Document(str(files[1]))
orig_title = orig.paragraphs[4].text
rev_title = rev.paragraphs[4].text
orig_date = orig.paragraphs[24].text
rev_date = next((p.text for p in rev.paragraphs if "October" in p.text and "2026" in p.text), "")
lines.append(f"ORIGINAL TITLE {orig_title!r} DATE {orig_date!r}")
lines.append(f"REVISED TITLE {rev_title!r}")
# role counts
for label, doc in (("ORIG", orig), ("REV", rev)):
    t = "\n".join(p.text for p in doc.paragraphs)
    lines.append(f"{label} 28-borrowers={t.count('28 borrowers')} 30-borrowers={t.count('30 borrowers')} three-collectors={t.count('three collectors')}")
out.write_text("\n".join(lines), encoding="utf-8")
print("ok")
