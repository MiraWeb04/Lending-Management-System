from pathlib import Path
from docx import Document

p = Path(r"C:\Users\HomePC\Downloads\FRAMS_REVISED_MANUSCRIPT_FINAL.docx")
d = Document(str(p))
out = Path(r"C:\xampp\htdocs\lending_system\_frams_revision\final_tables.txt")
with out.open("w", encoding="utf-8") as f:
    for ti in [10, 12, 17, 20]:
        t = d.tables[ti]
        f.write(f"\n===== T{ti} {len(t.rows)}x{len(t.columns)} =====\n")
        for ri, row in enumerate(t.rows):
            f.write(f"R{ri}: {' || '.join(c.text.replace(chr(10),' / ')[:180] for c in row.cells)}\n")
    f.write("\n=== webhook paras ===\n")
    for i, para in enumerate(d.paragraphs):
        if "webhook" in para.text.lower() or "49" in para.text and "PayMongo" in para.text:
            f.write(f"[{i}] {para.text[:320]}\n")
    f.write("\n=== 12 PayMongo / twelve PayMongo / 49 of 49 ===\n")
    for i, para in enumerate(d.paragraphs):
        t = para.text
        if any(s in t for s in ["twelve PayMongo", "12 PayMongo", "49 of 49", "12/12 cases", "PM-01 to PM-12"]):
            f.write(f"[{i}] {t[:300]}\n")
print("ok")
