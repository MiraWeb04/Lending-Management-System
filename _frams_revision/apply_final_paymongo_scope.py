# -*- coding: utf-8 -*-
"""Remove webhook tests from official manuscript scope; keep polling + staff reconciliation as implemented workflow."""
from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

SOURCE = Path(r"C:\Users\HomePC\Downloads\FRAMS_REVISED_MANUSCRIPT.docx")
FINAL = Path(r"C:\Users\HomePC\Downloads\FRAMS_REVISED_MANUSCRIPT_FINAL.docx")
REPORT = Path(r"C:\Users\HomePC\Downloads\FRAMS_PAYMONGO_SCOPE_REPORT.docx")
WORKSPACE = Path(r"C:\xampp\htdocs\lending_system\_frams_revision")

# Official remaining PayMongo cases after excluding webhook-only PM-11.
PM_CASES = 11  # PM-01..PM-10 and PM-12
UNIT = 6
IT = 7
ST = 12
PT = 12
TOTAL = UNIT + IT + ST + PM_CASES + PT  # 48


def set_run_font(run, name="Times New Roman", size=12, bold=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def replace_text(paragraph: Paragraph, new_text: str) -> None:
    drawing_runs = []
    for run in paragraph.runs:
        xml = run._r.xml
        if "w:drawing" in xml or "w:pict" in xml or "w:object" in xml:
            drawing_runs.append(run._r)
    for run in list(paragraph.runs):
        run._r.getparent().remove(run._r)
    for el in drawing_runs:
        paragraph._p.append(el)
    run = paragraph.add_run(new_text)
    set_run_font(run)


def set_cell(cell, text: str) -> None:
    if not cell.paragraphs:
        return
    replace_text(cell.paragraphs[0], text)
    for extra in cell.paragraphs[1:]:
        extra.text = ""


def iter_paragraphs(doc: Document):
    for para in doc.paragraphs:
        yield para
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    yield para


def delete_table_row(table, row_idx: int) -> None:
    row = table.rows[row_idx]
    table._tbl.remove(row._tr)


EXACT = {
    "In addition to payments recorded at the lending office, FRAMS includes PayMongo Dynamic QR Ph. A collector starts the code for an assigned loan when the method is bank transfer or e-wallet. The borrower pays by scanning that code. The QR image is not a payment confirmation. While the QR payment page remains open, FRAMS polls the PayMongo Payment Intent API and records the repayment only after the provider status is succeeded and the paid amount matches the pending transaction. A webhook handler is present in the application and may also post a verified event, but it is optional for this collector workflow. Showing the code, returning to the browser, or waiting for a local timer does not, by itself, record the repayment. Cash is recorded directly by authorized staff. Bank transfer and e-wallet may also be recorded directly, without a QR code, by an administrator or the assigned collector.": (
        "In addition to payments recorded at the lending office, FRAMS includes PayMongo Dynamic QR Ph. A collector starts the code for an assigned loan when the method is bank transfer or e-wallet. The borrower pays by scanning that code. The QR image is not a payment confirmation. While the QR payment page remains open, FRAMS polls the PayMongo Payment Intent API every five seconds and records the repayment only after the provider status is succeeded and the paid amount matches the pending transaction. If that page is closed while the attempt is still pending, expired, or failed, an authorized administrator or the assigned collector can select Check with PayMongo on the payments page. That staff action retrieves the same Payment Intent status and records the repayment only when the amount matches. Showing the code, returning to the browser, or waiting for a local timer does not, by itself, record the repayment. Cash is recorded directly by authorized staff. Bank transfer and e-wallet may also be recorded directly, without a QR code, by an administrator or the assigned collector."
    ),
    "The system includes PayMongo Dynamic QR Ph for supported non-cash collection transactions. According to the implemented workflow, an authorized collector initiates a QR payment for an assigned loan. The corresponding repayment is recorded after the Payment Intent status is confirmed as succeeded and the paid amount has been matched to the pending transaction. While the QR page is open, FRAMS polls the Payment Intent API for that confirmation. A webhook may also post a verified event when configured, but the collector workflow does not require a delivered webhook in order to record a matching paid intent. Cash payments and other authorized payment records may be entered directly by authorized personnel. Successful QR Ph collection depends on valid PayMongo API credentials and network access to retrieve Payment Intent status.": (
        "The system includes PayMongo Dynamic QR Ph for supported non-cash collection transactions. According to the selected implementation, an authorized collector initiates a QR payment for an assigned loan. The corresponding repayment is recorded after the Payment Intent status is confirmed as succeeded and the paid amount has been matched to the pending transaction. While the QR page is open, FRAMS polls the Payment Intent API for that confirmation. If the QR page is no longer open, authorized staff may start the same confirmation from the payments page using Check with PayMongo. Cash payments and other authorized payment records may be entered directly by authorized personnel. Successful QR Ph collection depends on valid PayMongo API credentials and network access to retrieve Payment Intent status."
    ),
    "PayMongo Dynamic QR Ph. A collector-initiated QR payment for bank transfer or e-wallet. The repayment is recorded after PayMongo reports a succeeded Payment Intent for the matching amount, through status polling on the open QR page or through an optional webhook. Displaying the QR image is not confirmation of payment.": (
        "PayMongo Dynamic QR Ph. A collector-initiated QR payment for bank transfer or e-wallet. The repayment is recorded after PayMongo reports a succeeded Payment Intent for the matching amount, through status polling on the open QR page or through staff-initiated Check with PayMongo. Displaying the QR image is not confirmation of payment."
    ),
    "Payments and penalties. The system records a payment with amount, date, method, collector, optional reference, and a penalty portion. It rejects a payment that exceeds the remaining balance plus the penalty still due. It recomputes the balance and the installment status. For an unpaid installment past its due date, it computes a penalty of 1 percent, 2 percent, 3 percent, or 5 percent of the installment amount for daily, weekly, semi-monthly, and monthly loans. Cash, bank transfer, and e-wallet can be recorded at once by an administrator or the assigned collector. For bank transfer or e-wallet, a collector may instead start PayMongo Dynamic QR Ph. That repayment is stored only after the Payment Intent is confirmed as succeeded, either through status polling on the open QR page or through an optional verified webhook.": (
        "Payments and penalties. The system records a payment with amount, date, method, collector, optional reference, and a penalty portion. It rejects a payment that exceeds the remaining balance plus the penalty still due. It recomputes the balance and the installment status. For an unpaid installment past its due date, it computes a penalty of 1 percent, 2 percent, 3 percent, or 5 percent of the installment amount for daily, weekly, semi-monthly, and monthly loans. Cash, bank transfer, and e-wallet can be recorded at once by an administrator or the assigned collector. For bank transfer or e-wallet, a collector may instead start PayMongo Dynamic QR Ph. That repayment is stored only after the Payment Intent is confirmed as succeeded, either through status polling on the open QR page or through staff-initiated Check with PayMongo."
    ),
    "Figure 2 is the Level 1 data flow diagram, with processes 1.0 through 5.0. The borrower, the collector, and the administrator send data into those processes and receive the records that belong to their role. Process 4.0 is a payment recorded by the collector. An administrator can record cash, bank transfer, and e-wallet on the same payments store. A confirmed PayMongo Payment Intent, retrieved while the QR page is open, also writes that store, and it is not drawn as a separate process on this figure. An optional webhook may write the same store when it is configured.": (
        "Figure 2 is the Level 1 data flow diagram, with processes 1.0 through 5.0. The borrower, the collector, and the administrator send data into those processes and receive the records that belong to their role. Process 4.0 is a payment recorded by the collector. An administrator can record cash, bank transfer, and e-wallet on the same payments store. A confirmed PayMongo Payment Intent, retrieved while the QR page is open or through staff-initiated Check with PayMongo, also writes that store, and it is not drawn as a separate process on this figure."
    ),
    "paymongo_qrph_webhook_events. Stores optional PayMongo webhook events for a QR Ph transaction and whether each event was processed. In the implemented collector workflow, the repayment is posted after Payment Intent status polling confirms success; the webhook is an additional path and is not required for the open QR page to record a matching paid intent.": (
        "paymongo_qrph_webhook_events. This table exists in the operational schema and stores provider event rows when they are present. It is not part of the selected payment-confirmation workflow or of the official PayMongo evaluation scope. In the selected workflow, a repayment is posted after Payment Intent status polling or staff-initiated Check with PayMongo confirms success and the amount matches."
    ),
    "Figure 5.7 is the immediate recording path for cash, bank transfer, and e-wallet. An administrator or the assigned collector can save that payment as soon as the amount is valid. PayMongo Dynamic QR Ph waits for later confirmation. The Create Dynamic QR Ph action is shown only for bank transfer or e-wallet, and only a collector can start it. The page stores a pending QR Ph transaction and shows the code. The balance stays unchanged until the Payment Intent is confirmed as succeeded and the paid amount matches the pending transaction. Displaying the QR code does not record the repayment. A failed, expired, or mismatched attempt does not post a payment. A verified QR Ph payment remains on the receipt record and is not deleted.": (
        "Figure 5.7 is the immediate recording path for cash, bank transfer, and e-wallet. An administrator or the assigned collector can save that payment as soon as the amount is valid. PayMongo Dynamic QR Ph waits for later confirmation. The Create Dynamic QR Ph action is shown only for bank transfer or e-wallet, and only a collector can start it. The page stores a pending QR Ph transaction and shows the code. While that page remains open, FRAMS polls the Payment Intent API. If the page is closed while the attempt is still pending, expired, or failed, authorized staff can select Check with PayMongo. The balance stays unchanged until the Payment Intent is confirmed as succeeded and the paid amount matches the pending transaction. Displaying the QR code does not record the repayment. A failed, expired, or mismatched attempt does not post a payment. A recorded QR Ph payment remains linked to the transaction."
    ),
}

FRAGMENT = [
    (
        "webhook-event handling, invalid-signature rejection, and protection of verified QR-linked repayments. The procedure used an isolated FRAMS test server and disposable test records to assess the specified payment workflows.",
        "unresolved pending payment attempts, receipt display, and payment-record listing. The procedure used an isolated FRAMS test server and disposable test records to assess the specified payment workflows. Payment Intent polling and staff-initiated Check with PayMongo are implemented on the payments page; they were not included as separately passed cases in Table 4.7 because no dedicated polling or staff-reconciliation screenshot series was supplied for those actions.",
    ),
    (
        "The testing evidence distinguishes provider rejection caused by an invalid test key from successful QR code creation. Local webhook simulations are documented as simulations and are not treated as evidence of actual external PayMongo webhook deliveries or live payment charges. The twelve PayMongo test cases, identified as PM-01 to PM-12, are presented in Table 4.7, with supporting evidence provided in Figures B26–B37.",
        "The testing evidence distinguishes provider rejection caused by an isolated invalid test key from successful QR code creation. The official PayMongo evaluation scope contains 11 test cases, PM-01 to PM-10 and PM-12, presented in Table 4.7, with supporting evidence in Figures B26–B35 and B37. Webhook signature validation (formerly PM-11, Figure B36) is excluded from the official count because webhook processing is outside the selected confirmation workflow. No live PayMongo charge is claimed.",
    ),
    (
        "The other refreshes a pending PayMongo QR Ph transaction by retrieving its Payment Intent status and returns the stored result. A QR Ph repayment is posted only after that status is succeeded and the amount matches, or after an optional verified webhook posts the same confirmation.",
        "The other refreshes a pending PayMongo QR Ph transaction by retrieving its Payment Intent status and returns the stored result. Authorized staff can start the same retrieval from the payments list by selecting Check with PayMongo. A QR Ph repayment is posted only after that status is succeeded and the amount matches.",
    ),
    (
        "The repayment, receipt, and balance change are saved only after Payment Intent polling on the open QR page, or an optional verified webhook, confirms success and the paid amount matches the pending transaction. An unverified, expired, or mismatched QR code does not post a payment, and a verified QR Ph payment is not deleted from the list.",
        "The repayment, receipt, and balance change are saved only after Payment Intent polling on the open QR page, or staff-initiated Check with PayMongo, confirms success and the paid amount matches the pending transaction. An unverified, expired, or mismatched QR code does not post a payment.",
    ),
    (
        "The PayMongo QR Ph testing documentation includes twelve test cases, PM-01 to PM-12, supported by the screenshots and recorded results in Figures B26–B37. These tests evaluated safe handling of rejected QR creation requests, payment amount validation, cash-versus-QR controls, collector assignment restrictions, authentication requirements, unresolved QR payment attempts, receipt and payment record display, webhook signature validation, and payment summaries. The evidence includes the safe rejection of QR creation after PayMongo returned HTTP 401 for the isolated invalid test key, as well as validation messages and restrictions demonstrated in the isolated FRAMS test environment. These results support the specific behaviors documented in the corresponding test cases.",
        "The official PayMongo QR Ph testing documentation includes 11 test cases, PM-01 to PM-10 and PM-12, supported by the screenshots and recorded results in Figures B26–B35 and B37. These tests evaluated safe handling of rejected QR creation requests, payment amount validation, cash-versus-QR controls, collector assignment restrictions, authentication requirements, unresolved QR payment attempts, receipt and payment-record display, and payment summaries. The evidence includes the safe rejection of QR creation after PayMongo returned HTTP 401 for the isolated invalid test key. Payment Intent polling and staff-initiated Check with PayMongo are described as implemented functions; they are not marked passed as separate cases in Table 4.7.",
    ),
    (
        "PM-09 and PM-10 document a locally simulated repayment and the display of its receipt and payment record within FRAMS. PM-11 documents local webhook signature-validation tests involving missing signatures, invalid signatures, and an expired timestamp. These are valid evidence of the tested local application behavior; however, they are not evidence of external PayMongo webhook delivery, successful live payment processing, or actual transfer of funds. The twelve documented PayMongo test cases and their corresponding evidence are presented in Table 4.7 and Figures B26–B37.",
        "PM-09 and PM-10 document receipt contents and the listing of a recorded E-Wallet repayment in the isolated test environment. They confirm payment recording and receipt access for that stored record. They do not demonstrate Payment Intent polling, staff-initiated Check with PayMongo, or a live PayMongo charge. Formerly numbered PM-11 (Figure B36) is excluded from Table 4.7 because it tested local webhook signatures, which are outside the selected confirmation workflow.",
    ),
    (
        "The testing summary reports 49 of 49 documented test cases as passed, comprising six unit tests, seven integration tests, twelve system tests, twelve PayMongo-related tests, and twelve penetration and security tests.",
        f"The testing summary reports {TOTAL} of {TOTAL} documented official test cases as passed, comprising six unit tests, seven integration tests, twelve system tests, 11 PayMongo-related tests, and twelve penetration and security tests.",
    ),
    (
        "Local PayMongo simulations do not independently demonstrate a successful external PayMongo charge or a provider-delivered webhook.",
        "The documented PayMongo cases do not independently demonstrate a successful live PayMongo charge, Payment Intent polling, or staff-initiated Check with PayMongo.",
    ),
    (
        "PayMongo QR Ph repayments are recorded after Payment Intent confirmation, not merely because a QR image was displayed.",
        "PayMongo QR Ph repayments are recorded after Payment Intent confirmation through status polling on the open QR page or staff-initiated Check with PayMongo, not merely because a QR image was displayed.",
    ),
    (
        "The PayMongo evidence for PM-01 to PM-12 is attached as Figures B26–B37. The test descriptions and results below are aligned to the behaviors shown in the supplied evidence. PM-01 and PM-02 document safe failure handling with an isolated invalid test key. PM-09 documents a local simulation and receipt display; it does not establish a live PayMongo charge or external webhook delivery.",
        "The official PayMongo evidence for PM-01 to PM-10 and PM-12 is attached as Figures B26–B35 and B37. The test descriptions and results below are aligned to the behaviors shown in the supplied evidence. PM-01 and PM-02 document safe failure handling with an isolated invalid test key. PM-09 documents receipt display for a recorded E-Wallet repayment in the isolated test environment; it does not establish Payment Intent polling, staff reconciliation, or a live PayMongo charge. Figure B36 is retained only as excluded local-endpoint evidence and is not counted.",
    ),
    (
        "Scenario: A test-mode simulation completes one pending QR on a disposable loan.",
        "Scenario: An authorized user opens the receipt for a recorded E-Wallet repayment in the isolated test environment.",
    ),
    (
        "Expected result: The transaction changes from pending to paid, one repayment and one receipt are stored, and the balance changes only after confirmation.",
        "Expected result: The receipt should display the recorded repayment details and the remaining balance after payment.",
    ),
    (
        "Actual result: This was a local signed webhook simulation, not a PayMongo delivery and not a live charge. The pending transaction changed to succeeded. One E-Wallet repayment of PHP 100.00 was stored with receipt RJRR-20261018-000004 and reference pay_frams_pm09_once. The receipt shows PHP 3,150.00 before payment and PHP 3,050.00 after payment.",
        "Actual result: The receipt showed RJRR-20261018-000004, E-Wallet, PHP 100.00 paid, PHP 3,150.00 before payment, and PHP 3,050.00 remaining. This check confirms receipt contents for a stored record in the isolated test environment. It does not demonstrate Payment Intent polling, staff-initiated Check with PayMongo, or a live PayMongo charge.",
    ),
    (
        "Scenario: Webhook requests are submitted without a signature, with an invalid signature, and with a timestamp older than five minutes to verify the endpoint's signature-validation mechanism.",
        "Scope note: Figure B36 is excluded from the official PayMongo evaluation. It shows local endpoint responses and is not a Payment Intent polling test, not a staff Check with PayMongo test, and not a live charge.",
    ),
    (
        "Expected result: Requests with a missing signature, an incorrect signature, or a timestamp older than five minutes should be rejected with HTTP 401 and return an invalid_signature status.",
        "This figure is not counted in Table 4.7.",
    ),
    (
        "Actual result: The local webhook response table showed HTTP 401 and invalid_signature for the missing-signature, incorrect-signature, and older-timestamp requests. The displayed evidence documents local signature-validation responses, not external PayMongo webhook deliveries.",
        "The screenshot is retained in the appendix file only so that the excluded evidence is not silently discarded. It is not treated as a passed official test case.",
    ),
    (
        "The documented functional test-case summary contains 49 cases: six unit tests, seven integration tests, twelve system tests, twelve PayMongo-related tests, and twelve penetration/security tests. The recorded case summary reports 49 passed and 0 failed. These counts describe the listed test cases and should not be interpreted as proof of exhaustive testing or as a count of every execution run.",
        f"The documented official functional test-case summary contains {TOTAL} cases: six unit tests, seven integration tests, twelve system tests, 11 PayMongo-related tests, and twelve penetration/security tests. The recorded official case summary reports {TOTAL} passed and 0 failed. These counts describe the listed official test cases and should not be interpreted as proof of exhaustive testing or as a count of every execution run.",
    ),
    (
        "The 49-case summary above remains the count of documented cases listed in the test tables.",
        f"The {TOTAL}-case summary above remains the count of documented official cases listed in the test tables.",
    ),
    (
        "The PayMongo and penetration/security evidence uses an isolated test server at 127.0.0.1:8099 and the frams_capstone_test database. The existing lending_management database was not modified, and no live PayMongo charge was created. PM-01 documents rejection of an isolated invalid test key, while PM-09 and PM-10 are local webhook simulations. These results support local validation and rejection-handling checks, but they do not prove a successful end-to-end transaction delivered by PayMongo’s external sandbox or production service. System tests ST-01 through ST-12 are recorded in Figures B14–B25.",
        "The PayMongo and penetration/security evidence uses an isolated test server at 127.0.0.1:8099 and the frams_capstone_test database. The existing lending_management database was not modified, and no live PayMongo charge was created. PM-01 documents rejection of an isolated invalid test key. PM-09 and PM-10 document receipt and payment-list display for a stored E-Wallet record. These results support local validation, rejection handling, and record display. They do not prove Payment Intent polling, staff-initiated Check with PayMongo, or a successful end-to-end live charge. System tests ST-01 through ST-12 are recorded in Figures B14–B25.",
    ),
    (
        "This appendix does not claim that every browser workflow or external PayMongo webhook was independently verified.",
        "This appendix does not claim that every browser workflow, Payment Intent polling sequence, staff Check with PayMongo action, or live PayMongo charge was independently verified.",
    ),
    (
        "This appendix does not contain live API keys, webhook secrets, or borrower records.",
        "This appendix does not contain live API keys, secrets, or borrower records.",
    ),
    (
        "it does not evidence duplicate-webhook handling.",
        "it does not demonstrate Payment Intent polling or staff-initiated Check with PayMongo.",
    ),
    (
        "does not evidence duplicate-webhook handling",
        "does not demonstrate Payment Intent polling or staff-initiated Check with PayMongo",
    ),
    (
        "Local PayMongo simulations are distinguished from external provider deliveries.",
        "Documented PayMongo cases are distinguished from live provider charges and from untested polling and staff-reconciliation actions.",
    ),
    (
        "PayMongo QR Ph repayments should be treated as confirmed transactions only when the required payment-verification process has been completed successfully.",
        "PayMongo QR Ph repayments should be treated as confirmed transactions only after Payment Intent status is succeeded and the paid amount matches, through polling on the open QR page or through staff-initiated Check with PayMongo.",
    ),
    (
        "Figure B36 — PM-11",
        "Figure B36 — excluded from official PayMongo scope",
    ),
]


def update_pm_tables(doc: Document) -> list[str]:
    notes = []
    for ti, table in enumerate(doc.tables):
        if len(table.rows) < 2:
            continue
        header0 = table.rows[0].cells[0].text.strip()
        first_id = table.rows[1].cells[0].text.strip() if len(table.rows) > 1 else ""
        # Table 4.7 / appendix PayMongo table
        if header0 in {"Test Case ID", "Test ID"} and first_id == "PM-01":
            # rewrite PM-10 remarks if webhook remains
            for row in table.rows:
                rid = row.cells[0].text.strip()
                if rid == "PM-10":
                    remarks = row.cells[-1].text
                    if "webhook" in remarks.lower() or "Webhook" in remarks:
                        set_cell(
                            row.cells[-1],
                            "PASSED — Figure B35 supports display of the payment record only; it does not demonstrate Payment Intent polling or staff-initiated Check with PayMongo.",
                        )
                if rid == "PM-09":
                    actual = row.cells[4].text
                    if "webhook" in actual.lower():
                        set_cell(
                            row.cells[4],
                            "The receipt showed RJRR-20261018-000004, E-Wallet, PHP 100.00 paid, PHP 3,150.00 before payment, and PHP 3,050.00 remaining. The record was stored in the isolated test environment and is not a live PayMongo charge.",
                        )
            # delete PM-11 row
            pm11_idx = None
            for ri, row in enumerate(table.rows):
                if row.cells[0].text.strip() == "PM-11":
                    pm11_idx = ri
                    break
            if pm11_idx is not None:
                delete_table_row(table, pm11_idx)
                notes.append(f"Table index {ti}: removed PM-11 row; official PayMongo rows now {len(table.rows)-1}.")
            else:
                notes.append(f"Table index {ti}: PM-01 table found, PM-11 already absent.")
        # Summary tables
        if header0 in {"Testing Level", "Testing level"}:
            for row in table.rows:
                label = row.cells[0].text.strip().lower()
                if "paymongo" in label:
                    set_cell(row.cells[1], str(PM_CASES))
                    set_cell(row.cells[2], str(PM_CASES))
                    if len(row.cells) > 3:
                        set_cell(row.cells[3], "0")
                    if len(row.cells) > 4:
                        # remarks or not-yet-tested
                        if "Remarks" in table.rows[0].cells[-1].text or table.rows[0].cells[-1].text.strip() == "Remarks":
                            set_cell(
                                row.cells[-1],
                                f"{PM_CASES}/{PM_CASES} official cases passed in the isolated test environment. PM-01/02 verified safe rejection with an invalid test key. PM-09/10 document receipt and payment-list display for a stored record. PM-11 (webhook signatures) is excluded. Polling and staff Check with PayMongo are implemented but not separately passed. See Figures B26–B35 and B37.",
                            )
                        else:
                            # T20 has not yet tested / blocked
                            set_cell(row.cells[4], "0")
                            set_cell(row.cells[5], "0")
                if label == "total":
                    set_cell(row.cells[1], str(TOTAL))
                    set_cell(row.cells[2], str(TOTAL))
                    if len(row.cells) > 3:
                        set_cell(row.cells[3], "0")
                    last = row.cells[-1].text
                    if "Remarks" in table.rows[0].cells[-1].text or "listed test" in last.lower() or "49" in last:
                        set_cell(
                            row.cells[-1] if table.rows[0].cells[-1].text.strip() == "Remarks" else row.cells[4] if len(row.cells) > 4 else row.cells[-1],
                            f"{TOTAL} listed official test cases passed: 6 unit, 7 integration, 12 system, {PM_CASES} PayMongo-related, and 12 penetration/security cases. Figure B36 is excluded from the official PayMongo count.",
                        )
            notes.append(f"Table index {ti}: summary counts set to PayMongo {PM_CASES} and total {TOTAL}.")
        # Entity table webhook_events purpose
        if header0 == "Entity":
            for row in table.rows:
                if row.cells[0].text.strip() == "paymongo_qrph_webhook_events":
                    set_cell(
                        row.cells[1],
                        "Schema table for provider event rows when present. Not part of the selected Payment Intent polling and staff Check with PayMongo confirmation workflow or of the official PayMongo test count.",
                    )
                    notes.append(f"Table index {ti}: paymongo_qrph_webhook_events purpose limited to schema documentation.")
    return notes


def write_report(path: Path, table_notes: list[str], remaining: list[str], remaining_detail: list[str]) -> None:
    d = Document()
    for section in d.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)
    h = d.add_paragraph("FRAMS PayMongo Scope Correction Report")
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in h.runs:
        set_run_font(run, size=16, bold=True)
    d.add_paragraph("23 October 2026")
    d.add_paragraph(
        "Source file: FRAMS_REVISED_MANUSCRIPT.docx. Original third revision FRAMS_3RDREVISIONS.docx was not overwritten. Corrected copy: FRAMS_REVISED_MANUSCRIPT_FINAL.docx."
    )

    ph = d.add_paragraph("Selected confirmation workflow")
    for run in ph.runs:
        set_run_font(run, size=13, bold=True)
    for b in [
        "PayMongo Payment Intent status polling on the open QR page (payments_lending.php?action=qrph_status, five-second interval).",
        "Staff-initiated Check with PayMongo on pending, expired, or failed QR rows (POST reconcile_qrph), which calls the same Payment Intent retrieval.",
        "QR image display is not payment confirmation.",
        "Webhook processing was removed from the official testing and evaluation scope. Source files and database structures were not deleted.",
    ]:
        d.add_paragraph(b, style="List Bullet")

    ph = d.add_paragraph("Official test counts after review")
    for run in ph.runs:
        set_run_font(run, size=13, bold=True)
    for b in [
        "Unit testing: 6 documented, 6 passed (unchanged).",
        "Integration testing: 7 documented, 7 passed (unchanged).",
        "System testing: 12 documented, 12 passed (unchanged).",
        "PayMongo testing: 12 documented cases reduced to 11 official cases by excluding PM-11. Remaining official IDs: PM-01 to PM-10 and PM-12. All 11 remain passed for their previously documented local conditions. No new passing case was added for polling or staff reconciliation.",
        "Penetration and security testing: 12 documented, 12 passed (unchanged).",
        f"Overall official total: {TOTAL} cases passed / {TOTAL} official cases (was 49/49).",
        "UAT means and respondent counts were not changed.",
    ]:
        d.add_paragraph(b, style="List Bullet")

    ph = d.add_paragraph("Why PM-11 was excluded and why PM-09/PM-10 were retained")
    for run in ph.runs:
        set_run_font(run, size=13, bold=True)
    for b in [
        "PM-11 tested local webhook signature rejection. That behavior is outside the selected Payment Intent polling and staff-reconciliation workflow, so it was removed from Table 4.7 and from the official count. Figure B36 remains in Appendix B with an exclusion note. It is not counted as passed.",
        "PM-09 and PM-10 are receipt and payment-list display checks for a stored E-Wallet record (RJRR-20261018-000004 / Payment ID 46). Those checks match the retained scope items for payment recording and receipt access. Their wording no longer treats webhook delivery as the confirmation method, and they are not treated as polling or staff-reconciliation passes.",
        "PM-12 remains a payment-summary display check and still does not test deletion.",
        "No dedicated screenshot series was supplied for the five-second Payment Intent poll or for the Check with PayMongo button, so those implemented functions are described in Chapters I, III, and IV but are not marked passed in Table 4.7.",
    ]:
        d.add_paragraph(b, style="List Bullet")

    ph = d.add_paragraph("Modified tables")
    for run in ph.runs:
        set_run_font(run, size=13, bold=True)
    for n in table_notes:
        d.add_paragraph(n, style="List Bullet")
    d.add_paragraph(
        "Chapter IV Table 4.7 and the Appendix B PayMongo table each lost the PM-11 row. Table 4.9 and the appendix testing-summary table now show PayMongo 11/11 and overall 48/48.",
        style="List Bullet",
    )

    ph = d.add_paragraph("Intentionally retained webhook-related references")
    for run in ph.runs:
        set_run_font(run, size=13, bold=True)
    for b in [
        "Entity-relationship discussion and Table of database entities: paymongo_qrph_webhook_events remains because it exists in the operational schema. The purpose text now states that it is not part of the selected confirmation workflow or official test count.",
        "Appendix L: the structure-only dump still lists paymongo_qrph_webhook_events and foreign key fk_paymongo_qrph_webhook_transaction. That is schema documentation, not a testing claim.",
        "Figure B36 filename and image remain in Appendix B, labeled as excluded local-endpoint evidence.",
        "Source code files includes/paymongo_qrph_webhook.php, the webhook HTTP endpoint, and related tables were not removed.",
        "No other webhook testing claims should remain in Chapters I–V official scope after this pass.",
    ]:
        d.add_paragraph(b, style="List Bullet")

    ph = d.add_paragraph("Final webhook-term search in FRAMS_REVISED_MANUSCRIPT_FINAL.docx")
    for run in ph.runs:
        set_run_font(run, size=13, bold=True)
    if remaining:
        d.add_paragraph("The following remaining hits contain the word webhook or a related identifier:")
        for item in remaining_detail:
            d.add_paragraph(item, style="List Bullet")
    else:
        d.add_paragraph("No remaining webhook-term hits were found.")

    d.save(path)


def main() -> None:
    shutil.copy2(SOURCE, FINAL)
    doc = Document(str(FINAL))

    for para in iter_paragraphs(doc):
        original = para.text
        if not original:
            continue
        new = original
        stripped = original.strip()
        if stripped in EXACT:
            new = EXACT[stripped]
        else:
            for old, repl in FRAGMENT:
                if old in new:
                    new = new.replace(old, repl)
        if new != original:
            replace_text(para, new)

    table_notes = update_pm_tables(doc)
    doc.save(str(FINAL))

    check = Document(str(FINAL))
    remaining_detail = []
    remaining_terms = []
    import re
    rx = re.compile(r"webhook|invalid_signature|evt_local|provider-delivered", re.I)
    for i, para in enumerate(check.paragraphs):
        if rx.search(para.text):
            remaining_detail.append(f"Paragraph {i}: {para.text[:240]}")
            remaining_terms.append("paragraph")
    for ti, table in enumerate(check.tables):
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                if rx.search(cell.text):
                    remaining_detail.append(f"Table {ti} R{ri}C{ci}: {cell.text[:240]}")
                    remaining_terms.append("table")

    write_report(REPORT, table_notes, remaining_terms, remaining_detail)
    WORKSPACE.mkdir(parents=True, exist_ok=True)
    shutil.copy2(FINAL, WORKSPACE / FINAL.name)
    shutil.copy2(REPORT, WORKSPACE / REPORT.name)

    print("FINAL paras", len(check.paragraphs), "tables", len(check.tables))
    print("REMAINING HITS", len(remaining_detail))
    for item in remaining_detail:
        print(item[:300])


if __name__ == "__main__":
    main()
