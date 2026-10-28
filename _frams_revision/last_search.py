import re
from pathlib import Path
from docx import Document

p = Path(r"C:\Users\HomePC\Downloads\FRAMS_REVISED_MANUSCRIPT_FINAL.docx")
d = Document(str(p))
rx = re.compile(r"webhook|invalid_signature|evt_local|PM-11|optional verified|provider-delivered", re.I)
out = Path(r"C:\xampp\htdocs\lending_system\_frams_revision\last_search.txt")
hits = []
for i, para in enumerate(d.paragraphs):
    if rx.search(para.text):
        hits.append(f"P{i}: {para.text.strip()}")
for ti, table in enumerate(d.tables):
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            if rx.search(cell.text):
                hits.append(f"T{ti}R{ri}C{ci}: {cell.text.strip()}")
# PM counts in summaries
counts = []
for ti in [10, 12, 17, 20]:
    t = d.tables[ti]
    counts.append(f"T{ti} {len(t.rows)}x{len(t.columns)} R4={ ' | '.join(c.text[:40] for c in t.rows[4].cells) }")
    if ti in (12, 20):
        counts.append("  TOTAL " + " | ".join(c.text[:50] for c in t.rows[6].cells))
out.write_text("HITS\n\n" + "\n\n".join(hits) + "\n\nCOUNTS\n" + "\n".join(counts), encoding="utf-8")
print("n_hits", len(hits))
