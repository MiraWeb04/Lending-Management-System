import docx

p = r"C:\Users\HomePC\Downloads\FRAMS_3RDREVISIONS.docx"
d = docx.Document(p)
out = r"C:\xampp\htdocs\lending_system\_frams_revision\table_ids.txt"

with open(out, "w", encoding="utf-8") as f:
    for ti in [0, 7, 8, 9, 10, 11, 12, 13]:
        t = d.tables[ti]
        f.write(f"\n===== TABLE {ti} {len(t.rows)}x{len(t.columns)} =====\n")
        for ri, row in enumerate(t.rows[:4]):
            cell0 = row.cells[0].text.replace("\n", "\\n")
            f.write(f"R{ri}C0={cell0!r} | R{ri}C1={row.cells[1].text[:80].replace(chr(10),' / ')!r}\n")
    f.write("\n=== para 177 full ===\n")
    f.write(d.paragraphs[177].text + "\n")
    f.write("\n=== para 181 full ===\n")
    f.write(d.paragraphs[181].text + "\n")
    f.write("\n=== para 32 full ===\n")
    f.write(d.paragraphs[32].text + "\n")
    f.write("\n=== para 68 full ===\n")
    f.write(d.paragraphs[68].text + "\n")
    f.write("\n=== para 582 full ===\n")
    f.write(d.paragraphs[582].text + "\n")
    f.write("\n=== para 623 full ===\n")
    f.write(d.paragraphs[623].text + "\n")
    f.write("\n=== para 842-845 ===\n")
    for i in range(842, 846):
        f.write(f"[{i}] {d.paragraphs[i].text}\n")
    f.write("\n=== body element types around appendix duplicate ===\n")
print("ok")
