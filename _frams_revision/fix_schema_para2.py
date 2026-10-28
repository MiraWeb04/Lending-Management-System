from pathlib import Path
import shutil
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

REVISED = Path(r"C:\Users\HomePC\Downloads\FRAMS_REVISED_MANUSCRIPT.docx")
WORKSPACE = Path(r"C:\xampp\htdocs\lending_system\_frams_revision") / REVISED.name

NEW = (
    "The operational database name in the connection file is lending_management, unless the environment supplies another name. "
    "The connection uses the host, database, user, and password from that configuration, with a local default. "
    "A structure-only schema dump of the operational database lending_management is provided as Appendix L and as the accompanying file FRAMS_APPENDIX_L_Database_Schema.sql. "
    "That dump includes the required tables, primary keys, unique keys, and the one declared foreign key. "
    "The PayMongo QR Ph tables are also created from the SQL files in database/migrations/. "
    "After the dump is imported, several feature columns continue to be created or extended when the related page first runs. "
    "Eligibility results are columns on the loan application, not a separate risk table. "
    "Passwords are stored as hashes. This section does not list credentials."
)

def replace_text(paragraph, new_text):
    for run in list(paragraph.runs):
        run._r.getparent().remove(run._r)
    run = paragraph.add_run(new_text)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(12)

doc = Document(str(REVISED))
for para in doc.paragraphs:
    if "A single installation script that creates every table was not found" in para.text:
        replace_text(para, NEW)
        break
else:
    raise SystemExit("paragraph not found")
doc.save(str(REVISED))
shutil.copy2(REVISED, WORKSPACE)
check = Document(str(REVISED))
print("remaining not found?", any("was not found in the repository" in p.text for p in check.paragraphs))
print("appendix L refs", sum("Appendix L" in p.text for p in check.paragraphs))
