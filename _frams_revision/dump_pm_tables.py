from pathlib import Path
from docx import Document

p = Path(r"C:\Users\HomePC\Downloads\FRAMS_REVISED_MANUSCRIPT.docx")
d = Document(str(p))
out = Path(r"C:\xampp\htdocs\lending_system\_frams_revision\pm_tables.txt")

with out.open("w", encoding="utf-8") as f:
    f.write("=== TABLE HEADS ===\n")
    for ti, t in enumerate(d.tables):
        head = " | ".join(c.text.replace("\n", " ")[:40] for c in t.rows[0].cells)
        f.write(f"T{ti} {len(t.rows)}x{len(t.columns)} {head}\n")

    f.write("\n=== PARAS 480-510 ===\n")
    for i in range(480, 511):
        t = d.paragraphs[i].text.strip()
        if t:
            f.write(f"[{i}] {t[:400]}\n")

    f.write("\n=== PARAS 650-720 ===\n")
    for i in range(650, 721):
        t = d.paragraphs[i].text.strip()
        if t:
            f.write(f"[{i}] {t[:400]}\n")

    f.write("\n=== PARAS 1020-1110 ===\n")
    for i in range(1020, 1111):
        t = d.paragraphs[i].text.strip()
        if t:
            f.write(f"[{i}] {t[:350]}\n")

    f.write("\n=== 49 / twelve / webhook remaining paras ===\n")
    for i, para in enumerate(d.paragraphs):
        t = para.text
        if any(s in t for s in ["49", "twelve PayMongo", "12 PayMongo", "PM-11", "webhook", "Webhook"]):
            f.write(f"[{i}] {t[:280]}\n")

print("ok")
