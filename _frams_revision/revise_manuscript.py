import re
import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

ROOT = Path(r"C:\xampp\htdocs\lending_system")
OUT = ROOT / "manuscript_revision"
OUT.mkdir(exist_ok=True)

WORKFLOW = (
    "An administrator, or a collector assigned to the loan, starts PayMongo Dynamic QR Ph "
    "for bank transfer or e-wallet. The system creates a Payment Intent and shows the QR code. "
    "While that page stays open, it asks PayMongo for the Payment Intent immediately and then every five seconds. "
    "A repayment, receipt, and balance update are saved only when the retrieved intent matches the stored amount, "
    "PHP currency, and payment mode, the intent status is succeeded, and a nested QR Ph payment is paid. "
    "If the page is closed first, an administrator or the assigned collector can check that one pending transaction "
    "from the Payments page. The same checks are used. The receipt is available only after those checks post the repayment. "
    "A second confirmation of the same provider payment does not post another repayment. "
    "Showing the QR code, a local timer, or a report that an e-wallet was charged does not record the payment. "
    "Pending transactions are not checked in the background."
)


def set_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def find_para(doc, start):
    matches = [p for p in doc.paragraphs if p.text.startswith(start)]
    if not matches:
        print("skip missing:", start[:80])
    return matches


def strip_until(doc, start_text, stop_text):
    removed = 0
    while True:
        paras = list(doc.paragraphs)
        start = next((i for i, p in enumerate(paras) if p.text.startswith(start_text)), None)
        if start is None:
            break
        stop = next((i for i, p in enumerate(paras) if i > start and p.text.startswith(stop_text)), None)
        if stop is None:
            raise SystemExit("Missing stop after " + start_text)
        for paragraph in paras[start:stop]:
            paragraph._element.getparent().remove(paragraph._element)
            removed += 1
    return removed


def replace_token(text, old, new):
    return re.sub(r"(?<![A-Za-z0-9])" + re.escape(old) + r"(?![A-Za-z0-9])", new, text)


def replace_everywhere(doc, old, new):
    count = 0
    for paragraph in doc.paragraphs:
        if old in paragraph.text:
            set_text(paragraph, replace_token(paragraph.text, old, new))
            count += 1
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if old in paragraph.text:
                        set_text(paragraph, replace_token(paragraph.text, old, new))
                        count += 1
    return count


def delete_paymongo_rows(table):
    rows = list(table.rows)
    ids = [row.cells[0].text.strip() for row in rows]
    for target in ("PM-11", "PM-10", "PM-09"):
        for row in list(table.rows):
            if row.cells[0].text.strip() == target:
                table._tbl.remove(row._tr)
                break
    for row in table.rows:
        if row.cells[0].text.strip() == "PM-12":
            set_text(row.cells[0].paragraphs[0], "PM-09")


def set_cell(cell, text):
    set_text(cell.paragraphs[0], text)
    for paragraph in cell.paragraphs[1:]:
        set_text(paragraph, "")


def revise_main(src, dest):
    shutil.copy(src, dest)
    doc = Document(dest)
    removed = strip_until(doc, "Figure B34 — PM-09", "Figure B37 — PM-12")
    for paragraph in list(doc.paragraphs):
        if paragraph.text.startswith("paymongo_qrph_webhook_events."):
            paragraph._element.getparent().remove(paragraph._element)
            removed += 1
    print(dest.name, "removed paragraphs", removed)

    replacements = {
        "In addition to payments recorded at the lending office": (
            "In addition to payments recorded at the lending office, FRAMS includes PayMongo Dynamic QR Ph. "
            + WORKFLOW
            + " Cash is recorded directly by authorized staff. Bank transfer and e-wallet may also be recorded directly, without a QR code, by an administrator or the assigned collector."
        ),
        "FRAMS includes PayMongo Dynamic QR Ph for bank transfer and e-wallet collection.": (
            "FRAMS includes PayMongo Dynamic QR Ph for bank transfer and e-wallet collection. "
            "An administrator, or a collector assigned to the loan, can start the QR code. "
            "A repayment is posted only after the server retrieves the Payment Intent and the amount, currency, mode, succeeded-intent, and paid QR Ph payment checks succeed. "
            "The integration uses the configured secret key, public key, and mode, which is test or live. "
            "Borrowers do not have a separate checkout page. The study does not include credit bureau integration, SMS integration, loan restructuring, deposits, savings, or investment accounts. "
            "Dependents and collateral notes may be stored, but they are not variables in the eligibility formula. The eligibility estimate remains rule-based and does not predict default. "
            "The study also does not establish that all recording errors have been eliminated or that the system has been deployed across multiple branches."
        ),
        "Payments and penalties.": (
            "Payments and penalties. The system records a payment with amount, date, method, collector, optional reference, and a penalty portion. "
            "It rejects a payment that exceeds the remaining balance plus the penalty still due. It recomputes the balance and the installment status. "
            "For an unpaid installment past its due date, it computes a penalty of 1 percent, 2 percent, 3 percent, or 5 percent of the installment amount for daily, weekly, semi-monthly, and monthly loans. "
            "Cash, bank transfer, and e-wallet can be recorded at once by an administrator or the assigned collector. "
            "For bank transfer or e-wallet, an administrator or the assigned collector may instead start PayMongo Dynamic QR Ph. "
            "That repayment is stored only after the Payment Intent checks succeed. The same provider payment is not recorded twice."
        ),
        "Figure 2 is the Level 1 data flow diagram": (
            "Figure 2 is the Level 1 data flow diagram, with processes 1.0 through 5.0. The borrower, the collector, and the administrator send data into those processes and receive the records that belong to their role. "
            "Process 4.0 is a payment recorded by the collector. An administrator can record cash, bank transfer, and e-wallet on the same payments store. "
            "A QR Ph repayment is written to that store only after the server confirms the Payment Intent. It is not drawn as a separate process on this figure."
        ),
        "Figure 5.7 is the immediate recording path": (
            "Figure 5.7 is the immediate recording path for cash, bank transfer, and e-wallet. An administrator or the assigned collector can save that payment as soon as the amount is valid. "
            "PayMongo Dynamic QR Ph waits for a later confirmation. The Create Dynamic QR Ph action is shown only for bank transfer or e-wallet. "
            "An administrator, or a collector assigned to the loan, can start it. The page stores a pending QR Ph transaction and shows the code. "
            "While the page is open, the server checks the Payment Intent every five seconds. The balance stays unchanged until the amount, currency, mode, succeeded-intent, and paid QR Ph payment checks succeed. "
            "If the page is closed, authorized staff can check that one pending transaction from the Payments page. A failed, expired, or mismatched attempt does not post a payment. "
            "A confirmed QR Ph payment remains on the receipt record and is not deleted."
        ),
        "PayMongo QR Ph testing covered": (
            "PayMongo QR Ph testing covered rejected QR creation with an isolated invalid test key, payment amount validation, cash-versus-QR controls, collector assignment, borrower and signed-out access restrictions, unresolved pending attempts, and protection of a QR-linked repayment from deletion. "
            "The test procedure used an isolated FRAMS test server and disposable test records. PM-01 and PM-02 document provider rejection, not a successful QR image. "
            "The retained cases and results are presented in Chapter IV and Figures B26–B34. Three earlier cases were removed from this official set because they are outside the implemented confirmation workflow."
        ),
        "PayMongo testing evaluated the integration": (
            "PayMongo testing evaluated QR Ph validation, access control, an unresolved pending attempt, and protection of a linked repayment. "
            "The nine retained PayMongo test cases, identified as PM-01 to PM-09, are presented in Table 4.7. "
            "They do not include a successful QR image or a completed wallet payment."
        ),
        "FRAMS was implemented as a browser-based application": (
            "FRAMS was implemented as a browser-based application for a local Apache, PHP, and MySQL environment. The documented local setup is XAMPP, with the application placed under the web server document root. "
            "Authorized users open the pages in a web browser. The server renders each page, checks the signed-in role, reads and writes the database, and returns HTML. "
            "There is no separate mobile application and no separate public web service for the screens. The payments page uses two authenticated JSON lookups. "
            "One loads the selected loan and its remaining balance into the payment form. The other returns the status of a PayMongo QR Ph transaction after the server retrieves the Payment Intent. "
            "A QR Ph repayment is posted only after the amount, currency, mode, succeeded-intent, and paid QR Ph payment checks succeed."
        ),
        "Payment and collection management.": (
            "Payment and collection management. The payments page lists collections and records a payment for an authorized loan. "
            "The form shows the borrower, the amount already paid, and the remaining balance. The fields stored are the amount, date, collector, method, optional reference number, and penalty portion. "
            "The receipt number uses the prefix RJRR, the date, and a numeric sequence. The borrower can view that borrower's payment history and download a receipt for that borrower's own payment. "
            "The collector works from assigned accounts. When the method is bank transfer or e-wallet, an administrator or the assigned collector can start PayMongo Dynamic QR Ph. "
            "The repayment, receipt, and balance change are saved only after the Payment Intent checks succeed. "
            "If the QR page has been closed, the Payments page lists that pending attempt and authorized staff can check it individually. "
            "An unverified, expired, or mismatched QR code does not post a payment, and a confirmed QR Ph payment is not deleted from the list. "
            "A payment that fails the amount checked in Section 4.1.1 is rejected. The same provider payment is not posted twice."
        ),
        "The results of unit, integration, system, PayMongo integration": (
            "The results of unit, integration, system, PayMongo integration, and security testing are organized in Tables 4.4 to 4.8 and Appendix B. "
            "The attached PayMongo evidence (Figures B26–B34) documents 9 retained checks in the isolated test environment, including safe rejection when the isolated test key was invalid. "
            "PM-01 and PM-02 are not successful QR generation. The attached penetration evidence (Figures B35–B46) documents 12 passed controlled security checks. "
            "The 12 system test cases remain Not Yet Tested in this report. Unit and integration testing entries are retained without changes. "
            "On 23 October 2026 a separate command-line check exercised the same server decision used by the five-second status check and by the Payments page check. "
            "It did not create a QR code and did not call PayMongo. An unpaid intent stayed pending, a failed payment was not posted, and a wrong amount stayed in review. "
            "An administrator could check a loan, a collector could check only an assigned loan, and a borrower or signed-out session could not. "
            "One accepted decision posted a single repayment, and repeating that provider payment did not post another. A cash entry still inserted its own payment row. "
            "That database transaction was rolled back, so the payment count, loan status, and temporary QR row were unchanged. "
            "The five-second timer was confirmed in the page source. It was not observed in a signed-in browser. An unsigned status request was redirected to staff sign-in. "
            "A separate wrong-currency trial was not executed in that check."
        ),
        "The PayMongo evidence for PM-01 to PM-12": (
            "The PayMongo evidence for PM-01 to PM-09 is attached after the main manuscript appendix as Figures B26–B34. "
            "All 9 retained checks passed in the isolated FRAMS test environment. "
            "PM-01 and PM-02 document safe failure handling with an isolated invalid test key, not successful QR generation. "
            "Provider messages and sensitive test-key material are redacted. "
            "This set does not show a completed wallet payment."
        ),
        "The literature reviewed in Chapter II": (
            "The literature reviewed in Chapter II places FRAMS alongside manual credit assessment and studies of predictive scoring. "
            "FRAMS uses configured eligibility rules and does not implement the machine-learning models discussed in some of the reviewed studies. "
            "The manuscript documents six unit test cases, seven integration test cases, twelve system test cases, nine PayMongo test cases, and twelve security test cases. "
            "These are documented cases, not proof that all forty-six cases passed. "
            "Final testing totals and user-evaluation means must be reconciled with the executed test records, evidence, and completed questionnaires in Appendix B."
        ),
        "All 49 functional test cases passed": (
            "The retained PayMongo evidence is 9 passed cases, and the penetration evidence is 12 passed cases. "
            "Unit and integration counts are unchanged and are not restated here as newly executed passes. "
            "The 12 system cases remain not yet tested. This paragraph does not claim that all 46 documented cases passed, and it does not claim that QR generation or a live payment passed."
        ),
        "The summary counts in this table show the updated outcomes": (
            "The summary counts in this table show the retained PayMongo and penetration outcomes and the existing not-yet-tested count for system testing. "
            "The total row marked with an asterisk covers only these updated categories; the Unit and Integration rows remain as previously documented and are not recalculated here. "
            "Therefore, this table must not be read as a claim that all 46 test cases passed."
        ),
        "PM-01 and PM-02 record the behavior available": (
            "PM-01 and PM-02 record the behavior available with the isolated invalid test key: PayMongo rejects creation, the QR page does not open, and the balance does not change. "
            "They do not show a PayMongo QR image."
        ),
    }
    for start, text in replacements.items():
        found = find_para(doc, start)
        for paragraph in found:
            set_text(paragraph, text)

    for table in doc.tables:
        header = " ".join(cell.text for cell in table.rows[0].cells)
        first_ids = [row.cells[0].text.strip() for row in table.rows]
        if "PM-09" in first_ids and "PM-12" in first_ids:
            delete_paymongo_rows(table)
        if table.rows[0].cells[0].text.strip() in ("Testing Level", "Testing level"):
            for row in table.rows:
                label = row.cells[0].text.strip()
                if label == "PayMongo Testing":
                    set_cell(row.cells[1], "9")
                    set_cell(row.cells[2], "9")
                    set_cell(row.cells[3], "0")
                    set_cell(row.cells[4], "9/9 retained cases passed in the isolated test environment. PM-01 and PM-02 verified safe rejection with an invalid test key, not successful QR generation. See Figures B26–B34.")
                elif label == "PayMongo testing":
                    set_cell(row.cells[1], "9")
                    set_cell(row.cells[2], "9")
                    set_cell(row.cells[3], "0")
                    set_cell(row.cells[4], "0")
                    if len(row.cells) > 5:
                        set_cell(row.cells[5], "0")
                elif label == "Total":
                    set_cell(row.cells[1], "46")
                    if len(row.cells) == 5:
                        set_cell(row.cells[2], "21*")
                        set_cell(row.cells[3], "0*")
                        set_cell(row.cells[4], "*Verified counts in the retained PayMongo and penetration evidence only. Unit and integration entries remain unchanged; 12 system cases remain not yet tested. Do not interpret this as all 46 cases passed.")
                    elif len(row.cells) == 6:
                        set_cell(row.cells[2], "21*")
                        set_cell(row.cells[3], "0*")
                        set_cell(row.cells[4], "12*")
                        set_cell(row.cells[5], "0*")
        for row in list(table.rows):
            joined = " ".join(cell.text for cell in row.cells)
            if "paymongo_qrph_webhook_events" in joined:
                table._tbl.remove(row._tr)

    replace_everywhere(doc, "PM-12", "PM-09")
    replace_everywhere(doc, "B37", "B34")

    def shift_later_figures(text):
        def repl(match):
            number = int(match.group(1))
            if 38 <= number <= 49:
                return "B" + str(number - 3)
            return match.group(0)
        return re.sub(r"B(\d{2})", repl, text)

    for paragraph in doc.paragraphs:
        updated = shift_later_figures(paragraph.text)
        if updated != paragraph.text:
            set_text(paragraph, updated)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    updated = shift_later_figures(paragraph.text)
                    if updated != paragraph.text:
                        set_text(paragraph, updated)

    doc.save(dest)
    return doc


def remaining_webhook(path):
    import zipfile
    xml = zipfile.ZipFile(path).read("word/document.xml").decode("utf-8", "ignore")
    text = re.sub(r"<[^>]+>", " ", xml)
    hits = []
    for match in re.finditer(r".{0,80}webhook.{0,80}", text, re.I):
        hits.append(" ".join(match.group(0).split()))
    return hits


main_src = ROOT / r"testing_evidence\_harness\FRAMS_REVISED_DRAFT_BEFORE_TEST_ALIGN.docx"
main_dest = OUT / "FRAMS_Manuscript_REVISED.docx"
revise_main(main_src, main_dest)
print("main webhook hits", len(remaining_webhook(main_dest)))
for hit in remaining_webhook(main_dest)[:20]:
    print(" HIT", hit[:200])

fig_src = ROOT / r"testing_evidence\FRAMS_QRPh_Figures_B26_B37_updated.docx"
fig_dest = OUT / "FRAMS_QRPh_Figures_REVISED.docx"
revise_main(fig_src, fig_dest)
print("figures webhook hits", len(remaining_webhook(fig_dest)))
for hit in remaining_webhook(fig_dest)[:20]:
    print(" FIG", hit[:200])
