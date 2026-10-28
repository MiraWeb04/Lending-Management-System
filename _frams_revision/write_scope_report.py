from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
import shutil

REPORT = Path(r"C:\Users\HomePC\Downloads\FRAMS_PAYMONGO_SCOPE_REPORT.docx")
WORKSPACE = Path(r"C:\xampp\htdocs\lending_system\_frams_revision")


def set_run_font(run, name="Times New Roman", size=12, bold=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


d = Document()
for section in d.sections:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

h = d.add_paragraph("FRAMS PayMongo Scope Correction Report")
h.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in h.runs:
    set_run_font(run, size=16, bold=True)
sub = d.add_paragraph("Corrected file: FRAMS_REVISED_MANUSCRIPT_FINAL.docx")
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
datep = d.add_paragraph("23 October 2026")
datep.alignment = WD_ALIGN_PARAGRAPH.CENTER

ph = d.add_paragraph("Files")
for run in ph.runs:
    set_run_font(run, size=13, bold=True)
for b in [
    "Original third revision left unchanged: FRAMS_3RDREVISIONS.docx.",
    "Working revised copy left unchanged: FRAMS_REVISED_MANUSCRIPT.docx.",
    "Corrected copy: FRAMS_REVISED_MANUSCRIPT_FINAL.docx.",
    "Webhook source files and database structures were not deleted.",
]:
    d.add_paragraph(b, style="List Bullet")

ph = d.add_paragraph("Selected confirmation workflow described in Chapters I, III, and IV")
for run in ph.runs:
    set_run_font(run, size=13, bold=True)
for b in [
    "Payment Intent API polling on the open QR page, every five seconds.",
    "Staff-initiated Check with PayMongo for pending, expired, or failed QR rows after the QR page is closed.",
    "Repayment posted only when Payment Intent status is succeeded and the paid amount matches.",
    "QR image display is not payment confirmation.",
    "Polling and staff Check with PayMongo are described as implemented. They are not marked passed in Table 4.7, because no dedicated screenshot series for those two actions was supplied.",
]:
    d.add_paragraph(b, style="List Bullet")

ph = d.add_paragraph("Official test counts")
for run in ph.runs:
    set_run_font(run, size=13, bold=True)
for b in [
    "Unit testing: 6 documented, 6 passed. Unchanged.",
    "Integration testing: 7 documented, 7 passed. Unchanged.",
    "System testing: 12 documented, 12 passed. Unchanged.",
    "PayMongo testing: official cases reduced from 12 to 11 by excluding PM-11. Remaining official IDs: PM-01, PM-02, PM-03, PM-04, PM-05, PM-06, PM-07, PM-08, PM-09, PM-10, and PM-12. All 11 remain passed for their previously documented local conditions. No new passing polling or reconciliation case was added.",
    "Penetration and security testing: 12 documented, 12 passed. Unchanged.",
    "Overall official total: 48 passed / 48 official cases. Previous official total was 49/49.",
    "UAT means, frequencies, and the 32-respondent workbook distribution were not changed.",
]:
    d.add_paragraph(b, style="List Bullet")

ph = d.add_paragraph("Modified tables")
for run in ph.runs:
    set_run_font(run, size=13, bold=True)
for b in [
    "Chapter IV Table 4.7 (document table index 10): removed the PM-11 row. Official PayMongo rows are now 11 (PM-01 to PM-10 and PM-12). PM-10 remarks no longer refer to duplicate-webhook handling.",
    "Appendix B PayMongo table (document table index 17): same PM-11 row removal and PM-10 remarks correction.",
    "Chapter IV Table 4.9 (document table index 12): PayMongo row 11/11; Total row 48/48. Remarks state that polling and staff Check with PayMongo are implemented but not separately passed, and that Figure B36 is excluded.",
    "Appendix testing-summary table (document table index 20): PayMongo 11/11/0/0/0; Total 48/48/0/0/0. The leftover 49-case remarks text in the Not yet tested column was cleared.",
    "Chapter III database-entity table (document table index 5): paymongo_qrph_webhook_events remains as a schema entity. Its purpose now states that it is not part of the selected confirmation workflow or official PayMongo count.",
]:
    d.add_paragraph(b, style="List Bullet")

ph = d.add_paragraph("Appendix B figure handling")
for run in ph.runs:
    set_run_font(run, size=13, bold=True)
for b in [
    "Figures B26–B35 and B37 remain the official PayMongo evidence set.",
    "Figure B36 remains in the file so the excluded screenshot is not silently discarded. It is labeled as excluded from official PayMongo scope and is not counted in Table 4.7.",
    "PM-09 and PM-10 remain as receipt and payment-list display checks for stored record RJRR-20261018-000004 / Payment ID 46. They are not treated as Payment Intent polling passes, staff-reconciliation passes, or live charges.",
]:
    d.add_paragraph(b, style="List Bullet")

ph = d.add_paragraph("Final document-wide search for webhook-related terms")
for run in ph.runs:
    set_run_font(run, size=13, bold=True)
d.add_paragraph(
    "After the correction pass, the only remaining hits for webhook or PM-11 in FRAMS_REVISED_MANUSCRIPT_FINAL.docx are the following intentionally retained references:"
)
for b in [
    "Chapter III ERD narrative: table name paymongo_qrph_webhook_events, identified as schema only and outside official evaluation scope.",
    "Chapter III entity table: entity name paymongo_qrph_webhook_events.",
    "Chapter IV testing discussion: one sentence noting that formerly numbered PM-11 / Figure B36 is excluded from Table 4.7.",
    "Appendix L: dump contents still list paymongo_qrph_webhook_events and foreign key paymongo_qrph_webhook_events.transaction_id, because those objects exist in the operational database. The dump does not claim they belong to the official test scope.",
    "Source code not altered: includes/paymongo_qrph_webhook.php still contains both the unused-for-scope webhook handler and the Payment Intent reconcile function used by polling and Check with PayMongo.",
]:
    d.add_paragraph(b, style="List Bullet")

ph = d.add_paragraph("What was not invented")
for run in ph.runs:
    set_run_font(run, size=13, bold=True)
for b in [
    "No new PayMongo pass/fail result was created.",
    "Polling was not marked passed.",
    "Staff Check with PayMongo was not marked passed.",
    "PM-11 was not replaced by a substitute passing case.",
    "UAT statistics were not recalculated or altered.",
]:
    d.add_paragraph(b, style="List Bullet")

d.save(str(REPORT))
shutil.copy2(REPORT, WORKSPACE / REPORT.name)
print("report saved")
