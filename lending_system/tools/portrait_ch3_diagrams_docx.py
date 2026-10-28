# -*- coding: utf-8 -*-
"""Place the Chapter 3 DFD and ERD on portrait pages."""
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from PIL import Image

SRC = Path(r"C:\Users\HomePC\Downloads\FRAMS_FINALDOCU_readable_diagrams.docx")
FALLBACK = Path(r"C:\Users\HomePC\Downloads\FRAMS_FINALDOCU_portrait.docx")
DFD = Path(r"C:\xampp\htdocs\lending_system\lending_system\documentation\frams_dfd_level1.png")
ERD = Path(r"C:\xampp\htdocs\lending_system\lending_system\documentation\frams_erd.png")

EMU = 914400
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
# A4 portrait text width is 8.27 - 1.5 - 1.0 = 5.77 in. Stay inside that.
IMG_W_IN = 5.60


def drawing_run(element):
    for child in element:
        if child.tag == f"{W}r" and child.find(f".//{W}drawing") is not None:
            return child
    return None


def clear_crop(run):
    drawing = run.find(f".//{W}drawing")
    for node in list(drawing.findall(".//" + A + "srcRect")):
        node.getparent().remove(node)


def set_extent(run, width_in, height_in):
    cx = str(int(width_in * EMU))
    cy = str(int(height_in * EMU))
    drawing = run.find(f".//{W}drawing")
    for ext in drawing.findall(".//" + qn("wp:extent")):
        ext.set("cx", cx)
        ext.set("cy", cy)
    for ext in drawing.findall(".//" + A + "ext"):
        if ext.getparent().tag.endswith("xfrm"):
            ext.set("cx", cx)
            ext.set("cy", cy)


def replace_blob(doc, run, png_path):
    blip = run.find(".//" + A + "blip")
    rid = blip.get(qn("r:embed"))
    doc.part.related_parts[rid]._blob = png_path.read_bytes()


def height_for(png_path):
    im = Image.open(png_path)
    return IMG_W_IN * im.size[1] / im.size[0]


def drop_section_breaks(doc):
    for paragraph in doc.paragraphs:
        pPr = paragraph._p.find(qn("w:pPr"))
        if pPr is None:
            continue
        for sect in list(pPr.findall(qn("w:sectPr"))):
            pPr.remove(sect)


def find_para(doc, snippet):
    hits = [p for p in doc.paragraphs if snippet in p.text]
    if len(hits) != 1:
        raise SystemExit(f"{snippet!r} matched {len(hits)} paragraphs")
    return hits[0]


def main():
    doc = Document(str(SRC))
    drop_section_breaks(doc)

    dfd_cap = find_para(doc, "Figure 2 Level 1 Data Flow Diagram of FRAMS")
    erd_cap = find_para(doc, "Figure 3 Entity-Relationship Diagram of FRAMS")
    dfd_run = drawing_run(dfd_cap._p.getprevious())
    erd_run = drawing_run(erd_cap._p.getprevious())
    if dfd_run is None or erd_run is None:
        raise SystemExit("Could not find the figure drawings")

    replace_blob(doc, dfd_run, DFD)
    replace_blob(doc, erd_run, ERD)
    clear_crop(dfd_run)
    clear_crop(erd_run)
    dfd_h = height_for(DFD)
    erd_h = height_for(ERD)
    set_extent(dfd_run, IMG_W_IN, dfd_h)
    set_extent(erd_run, IMG_W_IN, erd_h)

    body = doc.element.body.find(qn("w:sectPr"))
    pg = body.find(qn("w:pgSz"))
    if qn("w:orient") in pg.attrib:
        del pg.attrib[qn("w:orient")]
    pg.set(qn("w:w"), "11906")
    pg.set(qn("w:h"), "16838")

    try:
        doc.save(str(SRC))
        saved = SRC
    except PermissionError:
        doc.save(str(FALLBACK))
        saved = FALLBACK
    print("saved", saved)
    print(f"DFD {IMG_W_IN:.2f} x {dfd_h:.2f} in")
    print(f"ERD {IMG_W_IN:.2f} x {erd_h:.2f} in")


if __name__ == "__main__":
    main()
